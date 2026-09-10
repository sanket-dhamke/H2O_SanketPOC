"""Build GATEZO_User_Manual.docx from in-app labels (not invented copy)."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "GATEZO_User_Manual.docx"
LOGO = ROOT / "app" / "assets" / "brand" / "gatezo" / "options" / "D-combo.png"
WORDMARK = ROOT / "app" / "assets" / "brand" / "gatezo" / "wordmark.png"

NAVY = RGBColor(0x0B, 0x2B, 0x33)
TEAL = RGBColor(0x0B, 0x6E, 0x8F)
INK = RGBColor(0x1C, 0x2B, 0x33)
MUTED = RGBColor(0x5A, 0x6B, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xC5, 0xD9, 0xE2)
CREAM = "F4F1EA"
TEAL_HEX = "0B6E8F"
NAVY_HEX = "0B2B33"


def set_run(run, *, size=11, bold=False, color=INK, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), kwargs[edge].get("val", "single"))
            el.set(qn("w:sz"), kwargs[edge].get("sz", "4"))
            el.set(qn("w:color"), kwargs[edge].get("color", "C5D9E2"))
            tcBorders.append(el)
    tcPr.append(tcBorders)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("GATEZO User Manual  ·  ")
    set_run(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run2 = p.add_run()
    run2._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run3 = p.add_run()
    run3._r.append(instr)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run4 = p.add_run()
    run4._r.append(fld2)
    run5 = p.add_run("  ·  Confidential — for members & staff")
    set_run(run5, size=9, color=MUTED)
    for r in (run2, run3, run4):
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        r.font.name = "Calibri"


def add_toc(paragraph):
    run = paragraph.add_run()
    fldBegin = OxmlElement("w:fldChar")
    fldBegin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    fldSep = OxmlElement("w:fldChar")
    fldSep.set(qn("w:fldCharType"), "separate")
    fldEnd = OxmlElement("w:fldChar")
    fldEnd.set(qn("w:fldCharType"), "end")
    run._r.append(fldBegin)
    run._r.append(instr)
    run._r.append(fldSep)
    hint = paragraph.add_run("Right-click and choose Update Field to refresh this contents list in Word.")
    set_run(hint, size=10, color=MUTED)
    run._r.append(fldEnd)


class Manual:
    def __init__(self):
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(1.8)
        sec.bottom_margin = Cm(2.0)
        add_page_number(sec)
        styles = self.doc.styles
        styles["Normal"].font.name = "Calibri"
        styles["Normal"].font.size = Pt(11)
        styles["Normal"].font.color.rgb = INK
        styles["Normal"].paragraph_format.space_after = Pt(8)
        styles["Normal"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        for i, size, space_before, space_after in (
            (1, 20, 18, 8),
            (2, 14, 16, 6),
            (3, 12, 12, 4),
        ):
            st = styles[f"Heading {i}"]
            st.font.name = "Calibri"
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = NAVY if i == 1 else TEAL
            st.paragraph_format.space_before = Pt(space_before)
            st.paragraph_format.space_after = Pt(space_after)
            st.paragraph_format.keep_with_next = True

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text, *, size=11, bold=False, color=INK, space_after=8, center=False):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(0)
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        set_run(run, size=size, bold=bold, color=color)
        return para

    def rich(self, parts, *, space_after=8):
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(space_after)
        for text, kwargs in parts:
            run = para.add_run(text)
            set_run(run, **kwargs)
        return para

    def note(self, text, label="Note"):
        table = self.doc.add_table(rows=1, cols=1)
        table.autofit = True
        cell = table.cell(0, 0)
        shade(cell, "EAF4F7")
        set_cell_border(cell, top={"color": TEAL_HEX}, left={"color": TEAL_HEX, "sz": "16"},
                        bottom={"color": TEAL_HEX}, right={"color": TEAL_HEX})
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{label}.  ")
        set_run(r, size=10, bold=True, color=TEAL)
        r2 = p.add_run(text)
        set_run(r2, size=10, color=INK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def warn(self, text):
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        shade(cell, "FBE9E9")
        set_cell_border(cell, top={"color": "B44"}, left={"color": "B44", "sz": "16"},
                        bottom={"color": "B44"}, right={"color": "B44"})
        p = cell.paragraphs[0]
        r = p.add_run("Important.  ")
        set_run(r, size=10, bold=True, color=RGBColor(0xB4, 0x44, 0x44))
        r2 = p.add_run(text)
        set_run(r2, size=10, color=INK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def steps(self, items):
        for i, item in enumerate(items, 1):
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.4)
            para.paragraph_format.space_after = Pt(4)
            para.paragraph_format.space_before = Pt(0)
            n = para.add_run(f"{i}.  ")
            set_run(n, size=11, bold=True, color=TEAL)
            t = para.add_run(item)
            set_run(t, size=11, color=INK)

    def bullets(self, items):
        for item in items:
            para = self.doc.add_paragraph(style="List Bullet")
            para.clear()
            para.paragraph_format.space_after = Pt(3)
            run = para.add_run(item)
            set_run(run, size=11, color=INK)

    def ui(self, label):
        return f"“{label}”"

    def table(self, headers, rows, col_widths=None):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            shade(cell, NAVY_HEX)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
            run = cell.paragraphs[0].add_run(h)
            set_run(run, size=10, bold=True, color=WHITE)
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = table.rows[r_i + 1].cells[c_i]
                if r_i % 2 == 1:
                    shade(cell, CREAM)
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                cell.paragraphs[0].paragraph_format.space_before = Pt(2)
                run = cell.paragraphs[0].add_run(str(val))
                set_run(run, size=10, color=INK)
            prevent_row_split(table.rows[r_i + 1])
        prevent_row_split(table.rows[0])
        if col_widths:
            for row in table.rows:
                for i, w in enumerate(col_widths):
                    row.cells[i].width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(8)
        return table

    def feature(self, title, who, path, summary, steps, extra=None):
        self.h(title, 3)
        meta = self.doc.add_paragraph()
        meta.paragraph_format.space_after = Pt(6)
        for label, val in (("Who", who), ("Where", path)):
            r = meta.add_run(f"{label}: ")
            set_run(r, size=10, bold=True, color=TEAL)
            r2 = meta.add_run(f"{val}    ")
            set_run(r2, size=10, color=INK)
        self.p(summary, size=11, space_after=6)
        self.steps(steps)
        if extra:
            self.note(extra)

    def page(self):
        self.doc.add_page_break()

    def save(self):
        self.doc.save(OUT)
        return OUT


def cover(m: Manual):
    for _ in range(3):
        m.doc.add_paragraph().paragraph_format.space_after = Pt(6)
    logo = LOGO if LOGO.exists() else (WORDMARK if WORDMARK.exists() else None)
    if logo:
        p = m.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        run = p.add_run()
        run.add_picture(str(logo), width=Inches(2.4))
    else:
        m.p("GATEZO", size=32, bold=True, color=NAVY, center=True, space_after=4)
    m.p("USER MANUAL", size=22, bold=True, color=TEAL, center=True, space_after=10)
    m.p(
        "Step-by-step guide to every feature",
        size=14,
        color=INK,
        center=True,
        space_after=6,
    )
    m.p(
        "For residents, parents, gate staff, society / preschool admins, and the platform owner.",
        size=11,
        color=MUTED,
        center=True,
        space_after=18,
    )
    m.table(
        ["", ""],
        [
            ["Product", "GATEZO — one AI-first app for gated societies and preschools"],
            ["Platforms", "Android app (APK). The same app also runs in a web browser for training."],
            ["How to read this", "Button and screen names match the app exactly, in quotation marks."],
            ["Plans", "Some features need Prime or Platinum. Locked tiles say which plan is required."],
        ],
        col_widths=(1.6, 5.2),
    )
    m.p("Open the matching chapter for your login. You only see the tabs for your role.", size=11, center=True)
    m.page()


def toc_page(m: Manual):
    m.h("Contents")
    m.p(
        "In Microsoft Word: click inside the list below → right-click → Update Field → Update entire table.",
        size=10,
        color=MUTED,
    )
    p = m.doc.add_paragraph()
    add_toc(p)
    m.page()


def ch_welcome(m: Manual):
    m.h("1.  Welcome to GATEZO")
    m.p(
        "GATEZO is the app your society or preschool uses instead of a mix of WhatsApp groups, "
        "paper visitor registers, and spreadsheets for dues. One login covers money, the gate, "
        "community, and records."
    )
    m.p("What the app is for", size=12, bold=True, color=NAVY, space_after=4)
    m.bullets([
        "Money — maintenance or fees bills, UPI / card payments, cash receipts, dues and reminders.",
        "People at the gate — visitors, deliveries, pre-approved guests, vehicle QR (Platinum), staff attendance (preschool).",
        "Community — announcements, neighbours’ / parents’ board, community market, helpdesk, directory.",
        "Facilities — clubhouse / hall booking, home services, trusted helpers, SOS.",
        "Records — a payment history for every flat or student, plus reports and backups on higher plans.",
        "Help — Ask GATEZO on Home (Prime), a full Assistant chat, and Help & how-to.",
    ])
    m.note(
        "Your organisation is either a Society or a Preschool. The buttons change wording "
        "(Flat ↔ Student, Maintenance ↔ Fees, Resident ↔ Parent). The steps are the same. "
        "See chapter 3."
    )
    m.h("How this manual is organised", 2)
    m.bullets([
        "Chapters 5–6: install, sign in, profile — everyone.",
        "Chapter 7: Resident / Parent — every everyday feature.",
        "Chapter 8: Guard / gate desk.",
        "Chapter 9: Society / preschool Admin.",
        "Chapter 10: Preschool extras (My child, pickups, staff, student fees).",
        "Chapter 11: Platform owner (superadmin).",
        "Chapter 12: paying with UPI.",
        "Chapter 13: Ask GATEZO.",
        "Chapter 14: if something does not work.",
    ])


def ch_roles(m: Manual):
    m.h("2.  Who uses GATEZO (roles)")
    m.p(
        "You do not pick a role. The office creates your account (or you register with a join code "
        "and wait for approval). After Sign In, the bottom tabs match your role."
    )
    m.table(
        ["Role", "You will see this label", "Bottom tabs", "Main job"],
        [
            [
                "Resident / Parent",
                "Resident (society) or Parent (preschool)",
                "Home · Maintenance / Fees · Visitors · Community",
                "Pay dues, approve visitors, gate passes, community, bookings.",
            ],
            [
                "Guard",
                "Gate desk",
                "Home · Gate · Staff* · Community · Gate log",
                "Log visitors, admit gate-pass codes, respond to SOS. *Staff tab is preschool only.",
            ],
            [
                "Admin",
                "Society admin or Preschool admin",
                "Home · Maintenance / Fees · Members / Accounts · Staff* · Community · Gate log",
                "Bills, people, approvals, reports, the gate log.",
            ],
            [
                "Platform owner",
                "GATEZO Platform owner",
                "Overview · Societies",
                "All organisations, plans, branding, backups, marketplace moderation.",
            ],
        ],
    )
    m.note("Staff appears only when the organisation is a preschool.")


def ch_vocab(m: Manual):
    m.h("3.  Society vs preschool words")
    m.p("Same screens. Different labels.")
    m.table(
        ["In a society you see", "In a preschool you see"],
        [
            ["Flat / unit", "Student"],
            ["Resident", "Parent"],
            ["Maintenance", "Fees"],
            ["Wing", "Class (Nursery, Jr KG, Sr KG, Preschool)"],
            ["Members", "Accounts"],
            ["Society balance", "Fees balance"],
            ["Amenities / Book clubhouse", "Hall booking / Book hall"],
            ["Clubhouse / Party Hall", "Open Hall / Activity Room"],
            ["Neighbours’ board", "Parents’ board"],
            ["Remind unpaid residents", "Remind pending fees"],
            ["Manage members & flats", "Manage staff & students"],
        ],
    )
    m.p(
        "This manual uses society words first, then preschool in parentheses when it helps. "
        "If your Home folders say School and Help instead of At the gate, you are on a preschool login."
    )


def ch_plans(m: Manual):
    m.h("4.  Plans (Base, Prime, Platinum)")
    m.p(
        "The platform owner sets the plan on your society or preschool. If a shortcut is locked, "
        "the app shows the required plan (for example “Prime feature”) and “Ask your GATEZO owner to upgrade.”"
    )
    m.p("Your plan: Home (admin) → Operations → “Your plan”, or Members → “Plans & pricing”.", size=11)
    m.table(
        ["Plan", "What it is for", "Includes"],
        [
            [
                "Base",
                "Core operations to run digitally",
                "Visitor log; maintenance / fee billing and online pay; late-fee policy; maintenance heads; Buy & Sell; announcements; helpdesk; directory; basic reports and profile.",
            ],
            [
                "Prime",
                "Automation, AI & engagement",
                "Everything in Base, plus Ask GATEZO / Assistant, Gate pass, amenities / hall booking, WhatsApp + email reminders, rent management, automated monthly backup and wing-wise exports.",
            ],
            [
                "Platinum",
                "Hardware-integrated smart gate",
                "Everything in Prime, plus vehicle QR registry and printable windshield passes, gate scanners, real-time entry/exit, priority support and custom branding.",
            ],
        ],
    )
    m.note(
        "Buy & Sell (Community market) is available on Base. Ask GATEZO on Home needs Prime. "
        "Vehicle passes and Vehicle gate need Platinum."
    )


def ch_install(m: Manual):
    m.h("5.  Install the app and sign in")
    m.h("5.1  Install on Android", 2)
    m.steps([
        "Open the install link or QR code from your office (an Expo / GATEZO APK page).",
        "Download the .apk file.",
        "If Android asks, allow Install unknown apps / Install from this source for your browser or Files app.",
        "Open the file and tap Install.",
        "Open GATEZO. The splash and sign-in screen show the GATEZO name.",
    ])
    m.note(
        "If you already had an older GATEZO / society app, install the new APK over it "
        "(version code 32 or higher includes the latest phone fixes). You can keep the same login."
    )
    m.h("5.2  Sign in", 2)
    m.p("Screen title: “Sign in”. Tagline on a standard login: “Smart living, simplified”.")
    m.steps([
        "Enter Email and Password.",
        "Tap Show if you want to check the password, then Hide.",
        "Tap Sign In. The first sign-in after the server has been idle can take extra seconds.",
    ])
    m.p("If sign-in fails, these are the usual messages:", size=11, space_after=4)
    m.bullets([
        "Empty fields → “Please enter both your email and password.”",
        "Wrong email or password → a message that the email or password is not right.",
        "Inactive organisation → the society/preschool is currently inactive.",
        "Network → “Can't reach the GATEZO server…”",
        "New self-registered account → wait until an admin approves you.",
    ])
    m.p(
        "Accounts are created by your admin. The login screen says: "
        "“Accounts are created by your admin. Contact them if you can't sign in.”"
    )
    m.h("5.3  Advanced server settings", 2)
    m.p(
        "Only use this if your office told you to. On Sign in, open “Advanced server settings” "
        "and enter the Server URL they gave you. Phones in production should already point at the hosted GATEZO API — do not type localhost."
    )
    m.h("5.4  Forgot password", 2)
    m.steps([
        "On Sign in, tap Forgot password?",
        "On “Reset password”, enter your email and tap Send code.",
        "Check email for a 6-digit Reset code. (If the office has not connected email, a developer banner may show the code on screen.)",
        "Enter Reset code, New password, and Confirm new password. Password rule: at least 8 characters, 1 upper, 1 lower, 1 number.",
        "Tap Reset password. You should see “Password reset” — “Your password has been updated. You can now sign in.”",
    ])
    m.h("5.5  Register as a resident (join code)", 2)
    m.p("Use this only if the office shared a self-registration code. Screen title: “Register as a resident”.")
    m.steps([
        "On Sign in, tap Register.",
        "Enter Society join code (example format: GM7K2Q).",
        "Enter Your flat / unit number (example: A-1002).",
        "Enter Full name, Email, Phone (optional), Password, Confirm password.",
        "Tap Register.",
        "You should see “Request sent” — “You'll be able to sign in once an admin approves your account.”",
        "Ask the office to Approve you on Members.",
        "Then Sign In with the email and password you chose.",
    ])
    m.h("5.6  Change password (already signed in)", 2)
    m.steps([
        "Tap the profile icon (top right on Home).",
        "Under Account & security, tap Change password.",
        "Enter Current password, New password, Confirm new password.",
        "Tap Update.",
    ])
    m.h("5.7  Log out", 2)
    m.steps([
        "Open profile (or tap Logout on the Home header if shown).",
        "Tap Log out.",
        "Confirm “Are you sure you want to log out?”",
    ])


def ch_profile(m: Manual):
    m.h("6.  Profile, notifications and support")
    m.p("Tap the profile icon on Home. Sections: Account details, Notifications, My orders & payments (residents), Account & security, Support & feedback.")
    m.feature(
        "Profile settings",
        "Everyone",
        "Home → profile icon",
        "See your name, unit (flat/student), role, and a short GATEZO member ID.",
        [
            "Turn Push notifications on or off. On web you may see: “This preference applies to the GATEZO mobile app.”",
            "Residents: turn “Show my phone in directory” on if neighbours / parents may call you from Directory. Off = number hidden.",
            "Residents: open My orders & payments for a history of orders and payments.",
            "Support & feedback: call, email, or send feedback. Hours shown in-app: Mon–Sat, 9am – 7pm IST.",
        ],
        extra="In-app support contacts are the numbers configured for your deployment (profile → Support & feedback).",
    )


def ch_resident(m: Manual):
    m.h("7.  Resident / Parent — every feature")
    m.p(
        "Bottom tabs: Home · Maintenance (Fees) · Visitors · Community. "
        "Home folders: society — At the gate, Bills & bookings, Your society, Community. Preschool — School, Help."
    )

    m.h("7.1  Home dashboard", 2)
    m.feature(
        "Home snapshot",
        "Resident / Parent",
        "Tab: Home",
        "Greeting (Good morning / afternoon / evening), your name, badge (flat or role), Logout, profile. "
        "Charts cover the last 30 days of visitors and bill status. Pull down to refresh.",
        [
            "Read “Needs your attention” if anything is waiting (pending visitor, unpaid bill).",
            "Visitor donut segments: Approved, Pending, Rejected, Left at gate.",
            "Tap a folder to expand shortcuts.",
            "Use Ask GATEZO on this page if your plan includes it (Prime).",
            "Emergency SOS sits on Home — use only for a real emergency (see 7.3).",
        ],
        extra="If the summary does not load, pull down to refresh. The rest of Home still works.",
    )

    m.h("7.2  Ask GATEZO (Home)", 2)
    m.feature(
        "Ask GATEZO",
        "Resident / Parent (Prime plan)",
        "Home → box titled “Ask GATEZO”",
        "Type (or speak on the phone) a question about bills, visitors, or bookings. You get an answer from your society data, often with a button to open the right screen.",
        [
            "Type in the box. Placeholder: “Type a question about bills, visitors, or bookings”.",
            "On the phone, tap the microphone if you want to speak. Allow microphone permission if asked.",
            "Send / wait for the answer. Tap the action button if one appears (for example to pay or book).",
            "Tap Clear to start again.",
            "For a longer chat: Community → Assistant.",
        ],
        extra="If voice returns an error, type instead. Voice needs the live GATEZO server to accept transcription.",
    )
    m.p("Examples to try:", size=11, space_after=4)
    m.bullets([
        "How much maintenance do I owe?",
        "How do I book the clubhouse?",
        "Who visited my flat?",
        "How do I raise a complaint?",
        "What can you help with?",
    ])

    m.h("7.3  Emergency SOS", 2)
    m.feature(
        "Emergency SOS",
        "Resident / Parent (and other roles from Home)",
        "Home → Emergency SOS  ·  or Community stack",
        "One tap alerts guards, admins and nearby responders. Types: Medical, Security, Fire, Other.",
        [
            "Open Emergency SOS.",
            "Choose the type.",
            "Confirm “Raise … SOS?” and tap Send SOS.",
            "You should see “SOS sent”. Stay where you are if it is safe.",
            "Optional: set responder skills in SOS settings (Doctor, Nurse, First aid, Security-trained, Other) so you can be called when nearby.",
        ],
    )
    m.warn("SOS is only for real emergencies. It alerts guards, admins and nearby neighbours / parents immediately.")

    m.h("7.4  Visitors", 2)
    m.feature(
        "Approve a visitor at the gate",
        "Resident / Parent",
        "Tab: Visitors  ·  or Home → At the gate → Visitors at gate",
        "When a guard logs someone for your flat, you get a push notification. Title: “Visitors”. Subtitle: “Approve or review your gate entries”.",
        [
            "Open the Visitors tab (or tap the notification).",
            "Find the person whose status is Waiting. You can see photo and purpose when the guard attached them.",
            "Tap Approve, Reject, or Leave at gate.",
            "Statuses you will see: Waiting, Approved, Rejected, Left at gate.",
            "Pull down to refresh. Empty: “No visitors yet. Pull down to refresh.”",
        ],
    )

    m.h("7.5  Gate passes (pre-approve guests)", 2)
    m.feature(
        "Create a gate pass",
        "Resident / Parent (Prime)",
        "Visitors → Gate passes  ·  or Home → At the gate → Gate pass",
        "Pre-approve a guest or delivery so the guard can admit them with a code — no call required. Title: “Gate passes”. Subtitle: “Pre-approve guests & deliveries”.",
        [
            "Open Gate passes.",
            "Tap to pre-approve a guest or delivery.",
            "Choose type: Guest, Delivery, Cab, Service, Other.",
            "Enter Guest name (or Delivery / company name) and Contact number if you have it.",
            "Choose validity: 2 hours, Today, Tomorrow, or 3 days.",
            "Tap Create pass.",
            "Share the code with your guest (Share code).",
            "At the gate the guard taps Admit a pre-approved pass, enters the code, and Admit.",
            "To stop a pass early: open it → Cancel pass → confirm.",
        ],
        extra="Pass statuses: Active, Admitted, Expired, Cancelled. Empty: “No gate passes yet.”",
    )

    m.h("7.6  Vehicle passes", 2)
    m.feature(
        "Register a vehicle and print a gate QR",
        "Resident / Parent (Platinum)",
        "Visitors → Vehicle passes  ·  or Home → At the gate → Vehicle passes",
        "Title: “Vehicle passes”. Subtitle: “Register vehicles & print gate QR”. Types: Car, Bike / Scooter, Other.",
        [
            "Open Vehicle passes.",
            "Add the vehicle (registration number and type).",
            "Use Print gate pass. The print footer says “Vehicle Gate Pass” and “Stick this QR on the windshield…”.",
            "Use Regenerate QR if a sticker is lost. Activate or deactivate as needed.",
        ],
    )

    m.h("7.7  Pay maintenance / fees", 2)
    m.feature(
        "Pay a bill",
        "Resident / Parent",
        "Tab: Maintenance (or Fees)  ·  Home → Bills & bookings → Pay maintenance",
        "Title: “Maintenance” or “Fees”. Subtitle: “Bills, payments & receipts”. Banner explains that pay opens with UPI first.",
        [
            "Open the Maintenance / Fees tab. The header shows outstanding.",
            "Use Filter by month if you need an older bill (All months, or a specific month).",
            "On an unpaid bill, tap Pay (or Pay ₹… for an instalment / partial).",
            "On “Pay via UPI”, choose Google Pay, PhonePe, Paytm, UPI ID / QR, or Card / net banking. Or Cancel.",
            "Complete checkout (Razorpay). The merchant name should show GATEZO.",
            "Full payment: receipt modal “Payment Receipt” → Download / Print receipt → Close. The bill shows PAID.",
            "Partial: “Payment received” with the remaining balance. Pay again later for the rest.",
        ],
        extra="Empty: “No bills yet.” or “No bills for the selected month.” If payment fails you will see “Payment failed”. Cash payments are recorded by the office, not by you — those receipts show CASH.",
    )
    m.p("Receipt method labels you may see: “Cash (collected offline)”, “Razorpay (UPI/Card)”, “Test mode”.")

    m.h("7.8  Rent agreement (rented flats)", 2)
    m.feature(
        "Submit and track a lease",
        "Resident (society, Prime) — rented flats",
        "Maintenance tab → “Rent agreement — submit & track (rented flats)”  ·  screen “Rent agreements”",
        "Subtitle for residents: “Submit & track your lease”.",
        [
            "Open Rent agreements from the Maintenance tab link.",
            "Submit the agreement details the office asked for.",
            "Track status and expiry from the same screen. The office verifies leases and can alert before expiry.",
        ],
    )

    m.h("7.9  Transparency", 2)
    m.feature(
        "Where your money goes",
        "Resident / Parent",
        "Home → Your society → Transparency  ·  or Maintenance stack → Transparency",
        "Title: “Transparency”. Subtitle: “Where your money goes”. A 0–100 score from the books: payments in, expenses out, sealed on a tamper-evident ledger.",
        [
            "Open Transparency.",
            "Read the score and money in vs out.",
            "The score combines: ledger intact, itemised expenses, books updated in 45 days, dues collected, payout account on file, recent committee updates.",
            "This page does not take a payment. It is a view of the books.",
        ],
    )

    m.h("7.10  Sustainability (society)", 2)
    m.feature(
        "Green score and water use",
        "Resident (society)",
        "Home → Your society → Sustainability  ·  or Maintenance → Sustainability",
        "Title: “Sustainability”. Subtitle: “Green score & water metering”. This does not change what you pay in maintenance.",
        [
            "Open Sustainability.",
            "See your flat, the society median, and a top-savers list.",
            "Scoring is from monthly litre readings the office records (or imports by CSV), compared with the median.",
        ],
    )

    m.h("7.11  Book clubhouse / hall", 2)
    m.feature(
        "Request an amenity slot",
        "Resident / Parent (Prime)",
        "Community → Book clubhouse (or Book hall)  ·  Home → Bills & bookings → Book clubhouse",
        "Resident title: “Amenities” or “Hall booking”. Tabs: Book, My bookings.",
        [
            "Open Book clubhouse / Book hall.",
            "On Book, pick the facility, date, and a slot (Morning / Afternoon / Evening). Add notes if needed.",
            "Submit. You should see “Request sent” — the admin must approve.",
            "Watch My bookings. Statuses: Pending approval → Approved · pay to confirm → Booked & paid (or Declined / Cancelled).",
            "When it says Approved · pay to confirm, tap Pay and complete UPI / card like a bill.",
        ],
        extra="If nothing is listed: “No amenities are open for booking yet. Your admin can enable the clubhouse from their Amenities screen.” (Preschool wording talks about halls.)",
    )

    m.h("7.12  Community — announcements and board", 2)
    m.feature(
        "Read notices and post on the board",
        "Resident / Parent",
        "Tab: Community",
        "Title: “Community”. Subtitle: “Announcements & neighbours' board” (or parents' board). Tabs: Announcements, Board.",
        [
            "Open Community. Official notices are under Announcements. Pinned items are marked Pinned. Use the listen bar if you want the notice read aloud (English / हिंदी / मराठी where offered).",
            "Switch to Board. Categories: General, For sale, Question, Lost & found, Recommend.",
            "Tap + (New post) to share something. Empty board: “No posts yet. Be the first to share something!”",
            "Empty announcements: “No announcements yet.” Only admins create announcements.",
        ],
    )

    m.h("7.13  Home Services", 2)
    m.feature(
        "Book cleaning, AC, trades, movers",
        "Resident (society)",
        "Home → Community → Home services  ·  Community shortcut “Home Services”",
        "Title: “Home Services”. Subtitle: “Book trusted partners into your flat”. Tabs: Services, My bookings.",
        [
            "Open Home Services. Search: “Search cleaning, AC, painting…”.",
            "Pick a service. On the detail screen choose a package (Choose a package), date and slot.",
            "Tap Book visit. You should see that the booking was placed.",
            "Track it under My bookings. Instant jobs confirm themselves.",
        ],
        extra="Not shown on preschool Home. Home may also show a Home Services strip of offers.",
    )

    m.h("7.14  Community market (Buy & Sell)", 2)
    m.feature(
        "Buy, sell, borrow, skills, group-buy",
        "Resident / Parent",
        "Home → Community → Community market  ·  Community shortcut “Buy & Sell”",
        "Title: “Community market”. Subtitle: “Buy, sell, borrow, share skills & group-buy”. Tabs: Browse, My listings.",
        [
            "Open Community market. Filter kinds: All types, Buy & Sell, Borrow / Lend, Skills, Group buy.",
            "Filter categories such as Furniture, Electronics, Vehicles, Home Decor, Kids Items, Food, Services, Others.",
            "Search: “What are you looking for?”",
            "Open a listing for Product details. Use Chat with seller or WhatsApp if shown.",
            "To list: open My listings (or +), add photos, price and description.",
            "Owners and admins can delete a listing. The platform owner can moderate Buy & Sell across organisations.",
        ],
    )

    m.h("7.15  Trusted helpers", 2)
    m.feature(
        "Rated helpers whose ratings travel",
        "Resident / Parent",
        "Home → Community → Trusted helpers",
        "Title: “Trusted helpers”. Subtitle: “Ratings that follow the worker across societies”.",
        [
            "Open Trusted helpers.",
            "Open a helper’s Trust Passport.",
            "Rate 1–5 stars and optionally comment. You should see that the rating now travels with this worker.",
        ],
    )

    m.h("7.16  Helpdesk", 2)
    m.feature(
        "Raise a ticket",
        "Resident / Parent",
        "Community → Helpdesk  ·  Home folder Helpdesk",
        "Title: “Helpdesk”. Quick actions: Call security, Call office, Directory, New ticket.",
        [
            "Open Helpdesk.",
            "Tap + Raise a ticket / New ticket.",
            "Pick a category: Plumbing, Electrical, Housekeeping, Security, Billing, General, Other.",
            "Enter subject and description, then submit.",
            "Open the ticket later to read status: Open, In progress, Resolved — and any resolution note.",
            "Use Call security or Call office when you need a voice call instead of a ticket.",
        ],
        extra="Empty: “No tickets yet. Tap 'Raise a ticket' to report an issue.”",
    )

    m.h("7.17  Directory", 2)
    m.feature(
        "Find a neighbour or the office",
        "Resident / Parent",
        "Community → Directory",
        "Title: “Member directory”. Subtitle: “Reach someone in your society/preschool”.",
        [
            "Open Directory.",
            "Search by name or flat/student.",
            "Call or WhatsApp if a number is shown. Numbers appear only when that person enabled “Show my phone in directory”.",
        ],
    )

    m.h("7.18  Services & helplines", 2)
    m.feature(
        "Curated helplines",
        "Resident / Parent",
        "Community → Services & help",
        "Title: “Services & helplines”. A list of vendor / helpline contacts for your organisation.",
        [
            "Open Services & help from the Community shortcuts.",
            "Tap a contact to call when numbers are listed.",
        ],
    )

    m.h("7.19  AGM & voting (society)", 2)
    m.feature(
        "Vote on motions and read minutes",
        "Resident (society)",
        "Home → Your society → AGM & voting  ·  Community stack",
        "Title: “AGM & voting”. Subtitle: “Motions, e-voting, quorum & minutes”.",
        [
            "Open AGM & voting.",
            "Open a motion, read the text, and vote while voting is open.",
            "Read minutes of past meetings on the same screen.",
        ],
        extra="Not shown on preschool Home.",
    )

    m.h("7.20  Help & how-to and Assistant", 2)
    m.feature(
        "In-app how-to",
        "Everyone",
        "Community → Help & how-to  ·  or profile",
        "Title: “Help & how-to”. Short topics that match this manual (Home, Ask GATEZO, Visitors, bills, Transparency, Sustainability, Community, Home Services, clubhouse, Helpdesk, AGM, Buy & Sell).",
        [
            "Open Help & how-to and tap a topic.",
            "Or Community → Assistant for a full chat (“Ask anything about your society/school”) with starter chips.",
        ],
    )

    m.h("7.21  Preschool parent extras", 2)
    m.p("If your Home folders are School and Help, also follow chapter 10 (My child, pickup passes, pay fees).")


def ch_guard(m: Manual):
    m.h("8.  Guard / gate desk — every feature")
    m.p("Bottom tabs: Home · Gate · Staff (preschool) · Community · Gate log. Home folder: Gate desk.")

    m.h("8.1  Home", 2)
    m.p("Shortcuts: Log a new visitor, View gate log, Help & how-to. Preschool also: Child pickup, Staff attendance. Emergency SOS is for responding to alerts.")

    m.h("8.2  Log a new visitor", 2)
    m.feature(
        "New visitor",
        "Guard",
        "Tab: Gate  ·  Home → Log a new visitor",
        "Title: “New visitor”. Subtitle: “Log an entry at the gate”.",
        [
            "Open the Gate tab.",
            "Optional: tap “Dictate details (AI)”, speak name / phone / flat / purpose, then “Stop & fill form”. Allow microphone if asked.",
            "Take or attach a photo (allow camera if asked).",
            "Enter name, phone, vehicle if any.",
            "Pick the flat (or student).",
            "Pick purpose: Guest, Delivery, Cab, Service, Other. Preschool also has Pickup and Drop.",
            "Submit. Society: “Sent” — residents are notified to approve. Preschool: “Entry logged” — no resident approval step.",
        ],
        extra="Errors you may see: “Missing info”, “Microphone needed”, “Camera needed”.",
    )

    m.h("8.3  Admit a pre-approved pass", 2)
    m.feature(
        "Gate-pass code",
        "Guard",
        "Gate tab → “Admit a pre-approved pass”",
        "Residents share a 6-digit code from Gate passes (Prime).",
        [
            "On New visitor, tap Admit a pre-approved pass.",
            "Enter the code the guest shows you.",
            "Check the name and purpose.",
            "Tap Admit. You should see “Admitted” and that the resident / parent was notified.",
        ],
    )

    m.h("8.4  Gate log", 2)
    m.feature(
        "Today’s entries",
        "Guard / Admin",
        "Tab: Gate log (labelled Visitors in some builds)  ·  Home → View gate log",
        "Title: “Gate log”. Subtitle: “All society/preschool gate entries”.",
        [
            "Open Gate log.",
            "Scan statuses: Waiting, Approved, Rejected, Left at gate.",
            "Preschool: use Mark exit when the person leaves, if that button is shown.",
            "Pull down to refresh.",
        ],
    )

    m.h("8.5  Child pickup (preschool)", 2)
    m.feature(
        "Verify pickup pass",
        "Guard (preschool)",
        "Home → Child pickup  ·  screen “Child pickup”",
        "Subtitle: “Verify the pass, then log pickup/drop”.",
        [
            "Open Child pickup.",
            "Enter Pickup pass code → Verify.",
            "If Authorized, tap Log pickup or Log drop-off.",
            "You should see “Logged” and that the parent was notified.",
            "If you see “Pass not valid — do not release”, do not hand over the child. Call the office.",
        ],
    )

    m.h("8.6  Staff attendance (preschool)", 2)
    m.p("Same steps as admin — see 10.4. Guard tab: Staff.")

    m.h("8.7  SOS as a responder", 2)
    m.steps([
        "Open Emergency SOS from Home when an alert is active.",
        "Respond or resolve as the screen offers.",
        "Use helpline / ambulance quick dial if the office configured numbers.",
    ])

    m.h("8.8  Community", 2)
    m.p("Guards can read announcements, use Directory, Helpdesk contacts, Help & how-to, and Assistant. Posting on the board is for residents and admins.")


def ch_admin(m: Manual):
    m.h("9.  Admin — every feature")
    m.p(
        "Bottom tabs: Home · Maintenance / Fees · Members / Accounts · Staff (preschool) · Community · Gate log. "
        "Home folders: Run the society (or Run the school), Gate, Operations."
    )

    m.h("9.1  Home and Society / School manager", 2)
    m.feature(
        "AI co-pilot",
        "Admin",
        "Home → Run the society → Society manager (or School manager)",
        "Title: “Society manager” / “Preschool manager”. Subtitle: “Your AI co-pilot”. Drafts notices and shows insights.",
        [
            "Open Society manager / School manager.",
            "Read proactive insights. Jump to finance, helpdesk, or gate devices from the action chips if shown.",
            "Pick a draft: Monthly notice, Dues reminder, or Money summary.",
            "Edit the title and body. Tap Post as announcement when ready.",
            "Open Transparency from the manager if you want the books score.",
        ],
        extra="If drafting fails with “AI is off”, the server does not have AI configured. You can still type announcements yourself on Community.",
    )

    m.h("9.2  Finances / Fees tab", 2)
    m.p("Opens on “Finances” or “Fees”. Balance card: Society balance / Fees balance, with Collected and Pending.")

    m.feature(
        "Generate monthly bills",
        "Admin",
        "Finance tab → Generate bills",
        "Modal: “Generate monthly bills”. Bills every unit for the chosen month.",
        [
            "Tap Generate bills.",
            "Pick the billing month. Choose a flat amount per unit, or generate from Maintenance heads if you set those.",
            "Confirm. Residents see the bill on Maintenance / Fees, including a category breakdown when heads were used.",
        ],
    )
    m.feature(
        "Add an expense",
        "Admin",
        "Finance tab → Add expense",
        "Modal: “Add expense”. Labelled expenses feed Transparency and the balance.",
        [
            "Tap Add expense.",
            "Enter amount, label / category, and save.",
        ],
    )
    m.feature(
        "Remind unpaid people",
        "Admin (Prime for WhatsApp + email reminders)",
        "Finance tab → Remind unpaid residents / Remind pending fees",
        "Sends a nudge to everyone with a balance, including partial payments. Payment details are included when configured.",
        [
            "Tap Remind unpaid residents (or Remind pending fees).",
            "Wait until sending finishes. People with the app get a push; email / WhatsApp go out when those channels are configured.",
        ],
    )
    m.feature(
        "Record cash / view collections",
        "Admin",
        "Finance tab → Record cash / view collections  ·  screen “Collections”",
        "Subtitle: “Maintenance & cash entries”. Residents pay online; cash is recorded here.",
        [
            "Open Record cash / view collections.",
            "Find the unpaid bill. Tap Cash.",
            "On “Record cash payment”, enter Collected by and Collector phone (optional).",
            "Save. A cash receipt is stored (method: Cash (collected offline)).",
        ],
    )
    m.feature(
        "Payment audit (flat / student ledger)",
        "Admin",
        "Finance → Payment audit — full flat history  ·  or tap a unit under Flat-wise / Student-wise status",
        "Full bill and payment history for one unit, from the beginning.",
        [
            "On Finances, expand a wing / class, tap a unit — or open Payment audit.",
            "Read every bill and payment with running totals. Use this for disputes.",
        ],
    )
    m.feature(
        "Maintenance heads",
        "Admin (society)",
        "Finance → Maintenance heads (sinking fund, water…)",
        "Split the monthly figure into heads (security, water, sinking fund, common, and so on). Residents see the split on the bill.",
        [
            "Open Maintenance heads.",
            "Add or edit heads and amounts.",
            "Generate bills from heads next month so the breakdown is on each bill.",
        ],
    )
    m.feature(
        "Late fee policy",
        "Admin",
        "Finance → Late fee policy",
        "Late fees apply only after due date plus grace — not on months paid in advance.",
        [
            "Open Late fee policy.",
            "Turn on “Charge a late fee on overdue bills”.",
            "Choose Flat once, Per day, or % / month. Enter amount, grace days, and optional cap.",
            "Save.",
        ],
    )
    m.feature(
        "Amenities / hall as admin",
        "Admin (Prime)",
        "Finance → Amenities & clubhouse bookings (or Open hall bookings)",
        "Create facilities (Morning / Afternoon / Evening slots by default), set prices, approve or decline resident requests.",
        [
            "Open the amenities / hall screen.",
            "Tap + to add a facility. Edit slots and prices.",
            "On pending requests, Approve or decline. After Approve, the resident pays to confirm.",
            "Enable or disable a facility when it should not be booked.",
        ],
    )
    m.feature(
        "Reports & backup",
        "Admin (Prime for exports / automated backup)",
        "Finance → Reports & backup (wing-wise, PDF)  ·  preschool: Reports & backup (visitors, staff)",
        "Title: “Reports & backup”.",
        [
            "Open Reports & backup.",
            "Run the wing-wise / class-wise or visitor / staff export the screen offers (PDF / DOC as provided).",
            "Use email backup if offered on your plan.",
        ],
    )
    m.feature(
        "Vendor marketplace",
        "Admin (society)",
        "Finance → Vendor marketplace (premium)",
        "Title: “Vendor marketplace”. Subtitle: “Rent society premises to vendors”.",
        [
            "Open Vendor marketplace.",
            "Create a listing / booking for a vendor.",
            "Share a Razorpay pay link or Mark paid when money is received.",
        ],
    )
    m.feature(
        "Pay to GATEZO (subscription)",
        "Admin",
        "Finance → GATEZO subscription — Pay to GATEZO",
        "Title: “Pay to GATEZO”. Subtitle: “Platform subscription”. Pays the platform, not your society account.",
        [
            "Open Pay to GATEZO.",
            "Complete checkout the same way as other Razorpay payments.",
        ],
    )

    m.h("9.3  Members / Accounts", 2)
    m.feature(
        "Self-registration code",
        "Admin",
        "Members tab → Self-registration code",
        "Share a code so residents / parents can Register on the login screen.",
        [
            "Open Members (Accounts in preschool).",
            "On Self-registration code, tap Share (or generate New code if you need to rotate it).",
            "Send the code only to people who should join.",
            "When they register, they appear as pending. Tap Approve or Reject.",
        ],
    )
    m.feature(
        "Create an account yourself",
        "Admin",
        "Members → + New account  ·  screen “New account”",
        "Create Resident / Parent, Guard, or Admin with a temporary password.",
        [
            "Tap + New account.",
            "Choose role chips: Resident (or Parent), Guard, Admin.",
            "Enter name, email, phone, unit, and a temporary password (same strength rule: 8+ chars, upper, lower, number).",
            "Tap Create account. Share the email and temp password. Ask them to Change password after first Sign In.",
        ],
        extra="Empty members list: “No accounts yet. Tap New account.” Enable / Disable / Delete exist on each row. Disable stops login without deleting history.",
    )
    m.feature(
        "Flats / Students and bulk setup",
        "Admin",
        "Members → Flats / Students  ·  Bulk setup",
        "Add units one by one, or generate a wing/floor structure, or Import CSV.",
        [
            "Open Flats (Students).",
            "Add a unit, or tap Bulk setup.",
            "Generate structure: enter wings (e.g. A,B), floors, flats per floor, starting floor/unit → preview → create. You can rename rows in preview (soft cap around 600 editable rows).",
            "Or Import CSV. Columns include flatNo, block, ownerName, occupancy, rentMaintenanceAmount, memberName, email, phone, role.",
            "Mark occupancy (owner / rented) and optional rented-flat maintenance override.",
        ],
    )
    m.feature(
        "Bank account (where payments land)",
        "Admin",
        "Members → Bank  ·  screen “Bank account”",
        "Payout UPI / bank details shown to residents when they pay.",
        [
            "Open Bank account.",
            "Enter UPI and/or bank details.",
            "Save. Residents see “Payments go to …” and “Society UPI · …” on the bills tab.",
        ],
    )
    m.feature(
        "Rent agreements (admin)",
        "Admin (society, Prime)",
        "Members stack → Rent agreements",
        "Subtitle: “Verify tenant leases & track expiry”.",
        [
            "Open Rent agreements.",
            "Verify submitted leases, track expiry, and keep rented-flat maintenance overrides consistent with Flats.",
        ],
    )
    m.feature(
        "Gate scanners and vehicle gate",
        "Admin (Platinum)",
        "Members → Gate scanners  ·  Home → Gate → Vehicle gate",
        "Title: “Gate scanners” — “Connect entry-lane devices”. Vehicle QR registry lives with Vehicle passes.",
        [
            "Open Gate scanners to connect entry-lane devices if your site has them.",
            "Open Vehicle gate / Vehicle passes to manage the QR registry.",
        ],
    )
    m.feature(
        "Plans & pricing (view)",
        "Admin",
        "Home → Operations → Your plan  ·  Members → Plans & pricing",
        "See Base / Prime / Platinum feature lists. Changing the plan is done by the GATEZO platform owner on Societies.",
        [
            "Open Plans & pricing.",
            "Read what your current package includes. Ask the owner if you need an upgrade.",
        ],
    )

    m.h("9.4  Community as admin", 2)
    m.steps([
        "Announcements: on Community → Announcements tap + → “New announcement” (title, message, optional Pin to top).",
        "Delete an announcement if needed (confirm “Remove this announcement?”).",
        "Board: moderate / delete posts. You can also post.",
        "Helpdesk: open any ticket. Move Open → In progress → Resolved and add a resolution note.",
        "Home Services: tab Orders to see resident bookings.",
    ])

    m.h("9.5  Gate log", 2)
    m.p("Same Gate log as guards (chapter 8.4). Admins see all entries for the organisation.")

    m.h("9.6  Assets & AMC", 2)
    m.feature(
        "Assets and service due dates",
        "Admin",
        "Home → Operations → Assets & AMC (preschool: School assets)",
        "Title: “Assets & AMC”. Subtitle: “Lifts, pumps, DG, fire gear & service” (school: rooms, kits & service due dates).",
        [
            "Open Assets & AMC.",
            "Add an asset and service / AMC due dates.",
            "Use the list to see what is due.",
        ],
    )

    m.h("9.7  AGM as admin (society)", 2)
    m.steps([
        "Open AGM & voting from Operations or Finance stack.",
        "Create a motion, open voting, record quorum, add minutes.",
        "Residents vote from their AGM screen.",
    ])

    m.h("9.8  Rental compliance (society)", 2)
    m.feature(
        "Rental compliance",
        "Admin (society)",
        "Home → Operations → Rental compliance",
        "Title: “Rental compliance”. Subtitle: “Verification, checklists & deposits”.",
        [
            "Open Rental compliance.",
            "Work through verification, checklists and deposits as the screen lists them.",
        ],
    )

    m.h("9.9  Sustainability as admin (society)", 2)
    m.steps([
        "Open Sustainability from Operations or Finance.",
        "Record litres used this month per flat (tap + or import CSV).",
        "GATEZO compares each flat with the median. Residents see their own score.",
    ])


def ch_preschool(m: Manual):
    m.h("10.  Preschool extras")
    m.p("These screens appear when the organisation type is preschool. Society logins do not see them.")

    m.h("10.1  My child (parent)", 2)
    m.feature(
        "Today at school and pickup people",
        "Parent",
        "Home → School → My child",
        "Title: “My child”. Subtitle: “Today at school & pickup passes”.",
        [
            "Open My child.",
            "Read Today at school updates (Meal, Nap, Mood, Activity, Health, Note) posted by the office.",
            "Under Authorized to pick up, tap + Add person. Relation chips: Parent, Grandparent, Driver, Relative, Other.",
            "Share the pickup QR / pass with that person.",
            "Remove a person if they should no longer collect the child — their pass stops working.",
        ],
    )

    m.h("10.2  Pickups & updates (admin)", 2)
    m.feature(
        "Post a daily update",
        "Preschool admin",
        "Home → Run the school → Pickups & updates",
        "Title: “Pickups & updates”. Subtitle: “Post updates & manage pickup passes”.",
        [
            "Open Pickups & updates.",
            "Select the student.",
            "Tap + Post update. Choose Meal, Nap, Mood, Activity, Health, or Note.",
            "Save. The parent sees it on My child.",
            "Manage authorized pickup people from the same area when needed.",
        ],
    )

    m.h("10.3  Child pickup (guard)", 2)
    m.p("See 8.5. Always Verify the pass. Never release if the screen says the pass is not valid.")

    m.h("10.4  Staff attendance", 2)
    m.feature(
        "Check in / check out staff",
        "Preschool admin and guard",
        "Tab: Staff  ·  Home shortcuts Staff attendance",
        "Title: “Staff attendance”. Role chips: Teacher, Helper, Security, Admin, Other.",
        [
            "Open Staff.",
            "Enter Staff name, optional Phone, pick a role chip.",
            "Tap Check in.",
            "When they leave, Check out on their row.",
            "Empty: “No staff checked in today.” Admins can export this via Reports & backup (visitors, staff).",
        ],
    )

    m.h("10.5  Student fees", 2)
    m.feature(
        "Track and remind per student",
        "Preschool admin",
        "Finance → Student fees — track & remind  ·  screen “Fees — Students”",
        "Student-wise fee tracking with partial payments and optional per-student reminder dates.",
        [
            "Open Student fees — track & remind.",
            "Review paid / pending per student.",
            "Send reminders from Fees (Remind pending fees) or from the student-fees screen when offered.",
        ],
    )

    m.h("10.6  Pay fees (parent)", 2)
    m.p("Same as 7.7, on the Fees tab. Home shortcut: Pay fees.")


def ch_owner(m: Manual):
    m.h("11.  Platform owner (superadmin)")
    m.p("Bottom tabs: Overview · Societies. This login is not a society member login.")

    m.h("11.1  Overview", 2)
    m.feature(
        "GATEZO Platform",
        "Superadmin",
        "Tab: Overview",
        "Title: “GATEZO Platform”. Counts across organisations, month snapshot, shortcuts to moderate Buy & Sell and Backup & recovery.",
        [
            "Open Overview.",
            "Use the month selector for a snapshot across tenants.",
            "Open Moderate Buy & Sell to search, disable, or delete listings across societies.",
            "Open Backup & recovery for platform backups (run / download / email).",
        ],
    )

    m.h("11.2  Societies — onboard an organisation", 2)
    m.feature(
        "Onboard an organization",
        "Superadmin",
        "Societies tab → add / Onboard an organization",
        "Create a Society or Preschool (org type). Seed data in a demo environment is society-only; preschools are created here.",
        [
            "Open Societies.",
            "Open Onboard an organization.",
            "Enter organisation name, type (society or preschool), and the first admin details as the form asks.",
            "Save. Share the admin email and temporary password.",
            "Optionally bulk-import units if the onboard form / CSV path is shown.",
        ],
    )

    m.h("11.3  Per-organisation actions", 2)
    m.p("On each society / preschool card:")
    m.bullets([
        "Add admin — create another admin for that tenant.",
        "Plan — set Base / Prime / Platinum (this unlocks or locks features for every member of that tenant).",
        "Branding — name / background. “Members will see the new name/background the next time they open the app.”",
        "Activate / Deactivate — deactivated tenants cannot sign in (login explains the organisation is inactive).",
        "Open a tenant, then a flat/student, for Payment audit across the platform.",
        "Backup & recovery from the Societies stack as well as Overview.",
    ])


def ch_pay(m: Manual):
    m.h("12.  How payments work (UPI, card, cash)")
    m.p("Online collections use Razorpay. Cash is recorded only by admin.")
    m.table(
        ["What", "Who pays", "How"],
        [
            ["Maintenance / fees bill", "Resident / Parent", "Pay → Pay via UPI sheet → Google Pay, PhonePe, Paytm, UPI ID / QR, or Card / net banking."],
            ["Amenity / hall booking", "Resident / Parent after admin Approve", "Same UPI sheet, when status is Approved · pay to confirm."],
            ["Cash at the office", "Resident hands cash; admin records", "Collections → Cash → Collected by."],
            ["GATEZO subscription", "Society / preschool admin", "Pay to GATEZO."],
            ["Vendor / venue booking", "Vendor (or admin marks paid)", "Vendor marketplace pay link or Mark paid."],
            ["Home Services visit", "Resident books in-app", "Book visit — partner / offline fulfilment as configured."],
        ],
    )
    m.note(
        "The checkout name should be GATEZO. If an old phone build still shows another name, install the latest APK. "
        "UPI apps must be installed on the phone for Google Pay / PhonePe / Paytm buttons to open those apps."
    )


def ch_ask(m: Manual):
    m.h("13.  Ask GATEZO and Assistant — what to type")
    m.p("Prime plan. Home box and Community → Assistant share the same help. Live amounts (what you owe, who visited) come from your data, not from this PDF.")
    m.table(
        ["You can ask", "Opens / explains"],
        [
            ["How much maintenance do I owe? / How much is pending?", "Your outstanding bills"],
            ["How is the bill divided?", "Maintenance heads split"],
            ["Where can I get a receipt?", "Paid bills on Maintenance"],
            ["Who visited my flat?", "Visitors"],
            ["How do I approve a visitor? / make a gate pass?", "Visitors / Gate passes"],
            ["How do I book the clubhouse?", "Amenities"],
            ["How do I book AC service?", "Home Services"],
            ["How do I raise a complaint?", "Helpdesk"],
            ["What is the transparency score? / Where did money go?", "Transparency"],
            ["What is the green score?", "Sustainability"],
            ["How do I vote in the AGM?", "AGM & voting"],
            ["How do I sell something?", "Community market"],
            ["What can GATEZO do?", "Help overview"],
        ],
    )
    m.p("Guards can ask gate questions (for example who the last visitor was). Admins can ask about pending dues and tickets.")


def ch_tips(m: Manual):
    m.h("14.  Tips, limits and troubleshooting")
    m.table(
        ["What you notice", "What to do"],
        [
            ["First Sign In is slow", "The hosted server may be waking up after idle. Wait and retry once."],
            ["Can't reach the GATEZO server", "Check internet. Do not set Server URL to localhost on a phone."],
            ["Locked shortcut / Prime feature", "Your organisation’s plan does not include it. Admin: Your plan. Owner: Societies → Plan."],
            ["Voice / mic 404 or Request failed", "Type in Ask GATEZO instead, or wait until the office’s live API includes transcription."],
            ["My number is missing in Directory", "Profile → turn on Show my phone in directory."],
            ["Late fee looks wrong", "It applies only after due date + grace, never on months paid in advance."],
            ["Visitor notification", "Tap it to jump to Visitors."],
            ["Keyboard covers a form", "The screen should scroll the field up. Tap the field again if needed."],
            ["Need human help", "Profile → Support & feedback, or Helpdesk → Call office."],
        ],
    )
    m.h("What the app does in the background", 2)
    m.bullets([
        "Push notifications for visitors, tickets, announcements, gate passes, and payments (when enabled).",
        "Reminders over push, and email / WhatsApp when the office has those connected (Prime).",
        "Monthly backup email to admin when backups are on (Prime).",
        "Each society / preschool’s data is kept separate.",
    ])
    m.h("Android and iPhone", 2)
    m.p(
        "Today members typically install the Android APK. GATEZO is the same Expo app on iPhone; "
        "iPhone distribution needs Apple TestFlight or the App Store (Apple Developer enrolment). "
        "Until that is published, iPhone users can use the web app if the office shares the link."
    )


def ch_appendix(m: Manual):
    m.h("Appendix A — Feature map (quick find)")
    m.table(
        ["I want to…", "Role", "Go here"],
        [
            ["Pay my bill / get a receipt", "Resident / Parent", "Maintenance / Fees → Pay"],
            ["Approve a guest at the gate", "Resident / Parent", "Visitors"],
            ["Let a delivery in without a call", "Resident / Parent", "Visitors → Gate passes"],
            ["Book the hall", "Resident / Parent", "Community → Book clubhouse / Book hall"],
            ["Raise a complaint", "Resident / Parent", "Community → Helpdesk"],
            ["Sell a sofa", "Resident / Parent", "Community → Buy & Sell"],
            ["Book AC service", "Resident (society)", "Home → Home services"],
            ["See where money went", "Resident / Admin", "Home → Your society → Transparency"],
            ["Vote in the AGM", "Resident (society)", "Home → Your society → AGM & voting"],
            ["Log a visitor", "Guard", "Gate"],
            ["Admit a gate-pass code", "Guard", "Gate → Admit a pre-approved pass"],
            ["Release a child at pickup", "Guard (preschool)", "Child pickup → Verify"],
            ["Generate this month’s bills", "Admin", "Finance → Generate bills"],
            ["Record cash", "Admin", "Collections → Cash"],
            ["Approve a new member", "Admin", "Members → Approve"],
            ["Post a notice", "Admin", "Community → Announcements → +"],
            ["Add a preschool", "Owner", "Societies → Onboard an organization"],
            ["Upgrade a society to Prime", "Owner", "Societies → Plan"],
        ],
    )
    m.h("Appendix B — Demo / training logins", 2)
    m.warn(
        "Use these only on the official demo environment your trainer gives you. "
        "Change passwords on any live society. Do not use demo passwords on a production society that has real residents."
    )
    m.p("Demo password for seeded users (when the trainer has loaded seed data): Password123")
    m.table(
        ["Email", "Role", "Typical tenant"],
        [
            ["owner@h2o.com", "Platform owner", "All organisations"],
            ["admin@h2o.com", "Admin", "Demo society (e.g. Green Valley)"],
            ["guard@h2o.com", "Guard", "Same demo society"],
            ["resident@h2o.com", "Resident (example unit A-101)", "Same demo society"],
        ],
    )
    m.p(
        "Other seeded demo emails (same password) may include admin2@h2o.com, guard2@h2o.com, "
        "resident1b@h2o.com, resident2@h2o.com, resident3@h2o.com, and a second society such as "
        "admin@skyline.com / guard@skyline.com / resident@skyline.com. Preschool tenants are created by the owner, not by seed."
    )
    m.h("Appendix C — Document control", 2)
    m.table(
        ["Item", "Detail"],
        [
            ["Product name", "GATEZO"],
            ["Audience", "Members, gate staff, office admins, platform owner"],
            ["Matches app screens as of", "September 2026 (AI-first GATEZO build)"],
            ["In-app twin", "Community → Help & how-to  ·  Ask GATEZO"],
            ["Not in this manual", "Server deploy, EAS build, and investor financials — those are operations docs, not member how-to."],
        ],
    )
    m.p(
        "GATEZO — one AI-first app for gated societies and preschools.",
        size=11,
        bold=True,
        color=NAVY,
        center=True,
        space_after=4,
    )
    m.p("End of user manual.", size=10, color=MUTED, center=True)


def build():
    m = Manual()
    cover(m)
    toc_page(m)
    ch_welcome(m)
    ch_roles(m)
    ch_vocab(m)
    ch_plans(m)
    ch_install(m)
    ch_profile(m)
    ch_resident(m)
    ch_guard(m)
    ch_admin(m)
    ch_preschool(m)
    ch_owner(m)
    ch_pay(m)
    ch_ask(m)
    ch_tips(m)
    ch_appendix(m)
    return m.save()


if __name__ == "__main__":
    path = build()
    print(path)

