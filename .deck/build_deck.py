"""Rebuild the GATEZO investor deck around the live product.

Keeps the original narrative (team, pre-seed placeholders, Pune GTM) and
replaces the generic 'gate-only / visual placeholder' story with what the
demo actually does today.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(r"C:\Users\rajud\H2O_SanketPOC")
DECK = ROOT / ".deck"
ASSETS = ROOT / "app" / "assets"
SHOTS = DECK / "shots"
OUT = DECK / "GATEZO_Investor_Pitch_Deck_Latest.pptx"
OUT_ALT = DECK / "GATEZO_Investor_Pitch_Deck_Updated.pptx"
WA_DIR = Path(
    r"C:\Users\rajud\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm"
    r"\LocalState\sessions\0F9A65987131B5156AD30805FACE3F94D24DBDDD"
    r"\transfers\2026-37"
)
WA_OUT = WA_DIR / "GATEZO_Investor_Pitch_Deck_Enhanced.pptx"
WA_BACKUP = WA_DIR / "GATEZO_Investor_Pitch_Deck_Enhanced_backup.pptx"

# 16:9 widescreen — same canvas as the original Google Slides export.
W, H = 13.333, 7.50

NAVY = "0B3A49"
TEAL = "0B6E8F"
TEAL_DK = "094F66"
INK = "1C2B33"
MUTED = "5B6B73"
LINE = "D5DEE3"
BG = "F3F6F8"
WHITE = "FFFFFF"
GREEN = "1F8A4C"
AMBER = "C45C26"
SOFT = "E8F3F6"
SOFT_G = "E8F6EE"
SOFT_A = "F8EEE6"


def C(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def set_fill(shape, hex_color: str, alpha: float | None = None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = C(hex_color)
    if alpha is None:
        return
    solid = shape._element.spPr.find(qn("a:solidFill"))
    srgb = solid.find(qn("a:srgbClr"))
    # OOXML alpha: 100000 = fully opaque.
    srgb.append(
        parse_xml(
            f'<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            f' val="{int(alpha * 100000)}"/>'
        )
    )


def no_line(shape):
    shape.line.fill.background()


def rect(slide, l, t, w, h, color, *, alpha=None, line=None):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    set_fill(sh, color, alpha)
    if line:
        sh.line.color.rgb = C(line)
        sh.line.width = Emu(6350)
    else:
        no_line(sh)
    return sh


def round_rect(slide, l, t, w, h, color, *, adj=0.12, alpha=None, line=None):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    try:
        sh.adjustments[0] = adj
    except Exception:
        pass
    set_fill(sh, color, alpha)
    if line:
        sh.line.color.rgb = C(line)
        sh.line.width = Emu(6350)
    else:
        no_line(sh)
    return sh


def txt(
    slide,
    l,
    t,
    w,
    h,
    text,
    *,
    size=14,
    bold=False,
    color=INK,
    align="left",
    font="Calibri",
    anchor="top",
    italic=False,
):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.anchor = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }.get(anchor, MSO_ANCHOR.TOP)
    align_e = {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
    }[align]
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align_e
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.italic = italic
        p.font.color.rgb = C(color)
        p.space_after = Pt(2)
    return box


def notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def footer(slide, n, total=21):
    rect(slide, 0, 7.28, W, 0.22, NAVY)
    txt(slide, 0.35, 7.28, 6, 0.22, "GATEZO  ·  CONFIDENTIAL  ·  PRE-SEED",
        size=9, bold=True, color=WHITE, anchor="middle")
    txt(slide, 11.2, 7.28, 1.8, 0.22, f"{n}  /  {total}",
        size=9, bold=True, color=WHITE, align="right", anchor="middle")


def kicker(slide, l, t, label):
    txt(slide, l, t, 8, 0.28, label.upper(), size=11, bold=True, color=TEAL)


def heading(slide, l, t, w, title):
    txt(slide, l, t, w, 0.48, title, size=26, bold=True, color=NAVY)


def pic(slide, path, l, t, w, h):
    if not Path(path).exists():
        return None
    return slide.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))


def round_png(src: Path, dest: Path, radius=36):
    im = Image.open(src).convert("RGBA")
    w, h = im.size
    mask = Image.new("L", (w, h), 0)
    draw = ImageDraw.Draw(mask)
    draw.rounded_rectangle((0, 0, w - 1, h - 1), radius=radius, fill=255)
    out = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    out.paste(im, (0, 0))
    out.putalpha(mask)
    dest.parent.mkdir(parents=True, exist_ok=True)
    out.save(dest)
    return dest


def phone_frame(src: Path, dest: Path):
    """Put the live home screenshot in a simple dark bezel."""
    im = Image.open(src).convert("RGBA")
    # Trim any accidental letterboxing.
    bg = im.convert("RGB")
    w, h = bg.size
    px = bg.load()

    def col_lit(x):
        for y in range(0, h, 6):
            r, g, b = px[x, y]
            if r + g + b > 40:
                return True
        return False

    left, right = 0, w
    for x in range(w):
        if col_lit(x):
            left = max(0, x)
            break
    for x in range(w - 1, -1, -1):
        if col_lit(x):
            right = min(w, x + 1)
            break
    im = im.crop((left, 0, right, h))
    w, h = im.size
    bezel = 22
    canvas = Image.new("RGBA", (w + bezel * 2, h + bezel * 2), (11, 58, 73, 255))
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, w - 1, h - 1), radius=28, fill=255)
    rounded = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    rounded.paste(im, (0, 0))
    rounded.putalpha(mask)
    # Outer rounding
    outer = Image.new("L", canvas.size, 0)
    ImageDraw.Draw(outer).rounded_rectangle(
        (0, 0, canvas.size[0] - 1, canvas.size[1] - 1), radius=40, fill=255
    )
    framed = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    framed.paste(canvas, (0, 0))
    framed.paste(rounded, (bezel, bezel), rounded)
    framed.putalpha(outer)
    dest.parent.mkdir(parents=True, exist_ok=True)
    framed.save(dest)
    return dest


def fit_round_top(src: Path, dest: Path, w=640, h=360, radius=36):
    """Photo that sits on a rounded card — round the top, square the bottom."""
    im = ImageOps.fit(Image.open(src).convert("RGB"), (w, h), Image.Resampling.LANCZOS)
    im = im.convert("RGBA")
    mask = Image.new("L", (w, h), 0)
    d = ImageDraw.Draw(mask)
    d.rounded_rectangle((0, 0, w - 1, h - 1), radius=radius, fill=255)
    d.rectangle((0, radius, w, h), fill=255)
    im.putalpha(mask)
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest)
    return dest


def crop_login_card(src: Path, dest: Path):
    """The older login capture is a stacked strip — keep the top complete card."""
    im = Image.open(src).convert("RGB")
    w, h = im.size
    # Top ~58% contains the full login card.
    crop = im.crop((0, 0, w, int(h * 0.58)))
    dest.parent.mkdir(parents=True, exist_ok=True)
    crop.save(dest)
    return dest


def prepare_media():
    media = DECK / "media"
    media.mkdir(exist_ok=True)
    home = SHOTS / "03-resident-home.png"
    if home.exists():
        phone_frame(home, media / "phone-home.png")
        round_png(home, media / "home-rounded.png", 28)
    ac = SHOTS / "04-ac-service.png"
    if ac.exists():
        round_png(ac, media / "ac-service.png", 24)
    login = SHOTS / "01-login.png"
    if login.exists():
        crop_login_card(login, media / "login-card.png")
        round_png(media / "login-card.png", media / "login-rounded.png", 28)
    login_wide = SHOTS / "01-login-wide.png"
    if login_wide.exists():
        round_png(login_wide, media / "login-wide.png", 20)
    roles = {
        "role-resident.png": ASSETS / "home-resident.jpg",
        "role-guard.png": ASSETS / "home-guard.jpg",
        "role-admin.png": ASSETS / "home-admin.jpg",
        "role-preschool.png": ASSETS / "preschool-bg.jpg",
    }
    for name, src in roles.items():
        if src.exists():
            fit_round_top(src, media / name, 640, 360, 40)
    svc_dir = ASSETS / "services"
    for fn in ("instant.jpg", "cleaning.jpg", "ac.jpg", "movers.jpg", "salon.jpg", "pest.jpg"):
        src = svc_dir / fn
        if src.exists():
            fit_round_top(src, media / f"svc-{src.stem}.png", 520, 320, 32)
    return media


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def slide_title(prs, media):
    s = blank(prs)
    hero = ASSETS / "society-bg-wide.jpg"
    if hero.exists():
        pic(s, hero, 0, 0, W, H)
    else:
        rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, 0, W, H, NAVY, alpha=0.62)
    rect(s, 0, 0, 0.14, H, TEAL)

    wordmark = ASSETS / "brand" / "gatezo" / "wordmark.png"
    icon = ASSETS / "brand" / "gatezo" / "icon.png"
    if not icon.exists():
        icon = ASSETS / "icon.png"

    round_rect(s, 0.55, 0.38, 5.7, 1.08, WHITE, adj=0.14)
    if wordmark.exists():
        pic(s, wordmark, 0.72, 0.48, 5.36, 0.88)
    elif icon.exists():
        pic(s, icon, 0.7, 0.48, 0.88, 0.88)
        txt(s, 1.72, 0.58, 4.2, 0.7, "GATEZO", size=32, bold=True, color=NAVY)

    txt(s, 6.5, 0.48, 6.2, 0.28, "LIVE PRODUCT DEMO", size=13, bold=True, color="B8E0EA")
    txt(s, 6.5, 0.82, 6.2, 0.45, "PRE-SEED  ·  PUNE-LED PILOTS  ·  WORKING PRODUCT",
        size=12, bold=True, color="D7EEF4")
    txt(s, 0.55, 1.62, 12.2, 0.55,
        "One AI-first app for gated societies and preschools.",
        size=22, color="E8F3F6")
    txt(s, 0.55, 2.18, 12.2, 0.55,
        "Ask GATEZO, the gate, UPI bills or school fees, and home services — without stitching MyGate to Urban Company.",
        size=15, color="D5E4EA")

    # Two verticals — equal weight
    round_rect(s, 0.55, 3.4, 6.0, 1.55, NAVY, adj=0.1, alpha=0.55)
    txt(s, 0.75, 3.52, 5.6, 0.28, "GATED SOCIETIES", size=12, bold=True, color="9ED4E2")
    txt(s, 0.75, 3.82, 5.6, 0.95,
        "Flat  ·  Maintenance  ·  Clubhouse  ·  Resident\nVisitor log, UPI bills, helpdesk, Buy & Sell, home services.",
        size=13, color=WHITE)
    round_rect(s, 6.75, 3.4, 6.05, 1.55, TEAL, adj=0.1, alpha=0.72)
    txt(s, 6.95, 3.52, 5.7, 0.28, "PRESCHOOLS  —  FIRST-CLASS, NOT A SKIN", size=12, bold=True, color="D7F4FB")
    txt(s, 6.95, 3.82, 5.7, 0.95,
        "Student  ·  Fees  ·  Hall  ·  Parent  ·  Pickup at the gate\nNever says “flat”, “maintenance” or “clubhouse”. Same codebase.",
        size=13, color=WHITE)

    chips = [
        ("FOUR ROLES", "Resident / parent · guard · admin · superadmin"),
        ("ASK GATEZO", "Type or speak. Answers from live site data."),
        ("UPI IN-APP", "Razorpay · Google Pay · PhonePe · Paytm"),
        ("ONE CODEBASE", "Mobile + web. Society and preschool."),
    ]
    for i, (a, b) in enumerate(chips):
        x = 0.55 + (i * 3.1)
        round_rect(s, x, 5.15, 2.95, 1.0, NAVY, adj=0.14, alpha=0.58)
        txt(s, x + 0.12, 5.22, 2.7, 0.28, a, size=11, bold=True, color="B8E0EA")
        txt(s, x + 0.12, 5.5, 2.7, 0.55, b, size=12, color=WHITE)

    txt(s, 0.55, 6.35, 12, 0.32,
        "Udit Chothani  ·  Sanket Dhamke  ·  Aniket Chaudhari",
        size=14, bold=True, color=WHITE)
    txt(s, 0.55, 6.72, 12, 0.32,
        "Pre-revenue  ·  Product in active development  ·  Ready for pilot validation",
        size=13, color="C5D8DE")
    notes(s, "Open with the live demo after this slide if the room is technical. "
          "Do not quote revenue or site counts — still pre-revenue.")
    return s


def slide_team(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.32, "Who we are")
    heading(s, 0.5, 0.55, 10, "The team")
    txt(s, 0.5, 1.08, 12, 0.35,
        "Founder-led, with a working multi-role product — not a slide-ware prototype.",
        size=14, color=MUTED)

    people = [
        ("Udit Chothani", "Founder  |  Business, strategy & growth",
         "Experience across business, sales and customer-facing operations.\n"
         "Focused on practical products around real operational pain.\n"
         "Driving GATEZO from concept to pilot and scalable GTM."),
        ("Sanket Dhamke", "Founding team member  |  Engineering",
         "Building the live product: Ask GATEZO, roles, society + preschool vocabulary, payments and the mobile + web app."),
        ("Aniket Chaudhari", "Founding team member  |  Marketing & sales",
         "Go-to-market, customer conversations and early society / preschool outreach for Pune pilots."),
    ]
    for i, (name, role, bio) in enumerate(people):
        x = 0.45 + i * 4.2
        round_rect(s, x, 1.6, 4.0, 3.55, WHITE, adj=0.08, line=LINE)
        rect(s, x, 1.6, 4.0, 0.1, TEAL if i == 0 else NAVY)
        txt(s, x + 0.2, 1.85, 3.6, 0.4, name, size=18, bold=True, color=NAVY)
        txt(s, x + 0.2, 2.25, 3.6, 0.4, role, size=12, bold=True, color=TEAL)
        txt(s, x + 0.2, 2.75, 3.6, 2.1, bio, size=13, color=INK)

    round_rect(s, 0.45, 5.35, 6.15, 1.7, WHITE, adj=0.1, line=LINE)
    txt(s, 0.65, 5.5, 5.8, 0.3, "Advisors & mentors", size=14, bold=True, color=NAVY)
    txt(s, 0.65, 5.85, 5.8, 0.95,
        "Open. Seeking domain mentors across proptech, community operations, "
        "payments and enterprise sales — to accelerate Pune pilots.",
        size=13, color=INK)

    round_rect(s, 6.8, 5.35, 6.05, 1.7, SOFT, adj=0.1)
    txt(s, 7.0, 5.5, 5.7, 0.3, "How we present ourselves", size=14, bold=True, color=NAVY)
    txt(s, 7.0, 5.85, 5.7, 0.95,
        "We are pre-revenue. The product is real. This raise is to turn a working "
        "demo into paying sites, not to fund a slide-ware prototype.",
        size=13, color=INK)
    footer(s, 2)
    notes(s, "Udit is Founder. Sanket and Aniket are founding team members — not co-founders.")
    return s


def slide_problem(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "Why this exists")
    heading(s, 0.5, 0.52, 12, "The problem is not 'another visitor app'")
    txt(s, 0.5, 1.05, 12.2, 0.4,
        "Committees already pay for a gate app, then still open WhatsApp, Excel and Urban Company.",
        size=15, color=MUTED)

    cards = [
        ("01", "The category is loud and expensive",
         "MyGate, NoBrokerHood and ADDA own awareness. Many societies complain about price and ad-heavy home screens — not about missing a QR at the gate."),
        ("02", "Work is still fragmented",
         "Residents, guards and managers stitch registers, calls, Excel bills, UPI to the treasurer's personal ID, and a second app for a plumber."),
        ("03", "Nobody asks the software",
         "Paying a bill, logging a guest or booking a hall still means hunting through tabs. There is no AI layer that answers from the society's own data."),
        ("04", "Services sit in another company",
         "Urban Company wins the home. Community apps show banners. The resident should book AC service in the same place they approve the delivery."),
    ]
    for i, (num, title, body) in enumerate(cards):
        x = 0.45 + (i % 2) * 6.4
        y = 1.55 + (i // 2) * 2.15
        round_rect(s, x, y, 6.15, 2.0, WHITE, adj=0.08, line=LINE)
        txt(s, x + 0.22, y + 0.18, 1.0, 0.35, num, size=16, bold=True, color=TEAL)
        txt(s, x + 0.22, y + 0.55, 5.7, 0.4, title, size=16, bold=True, color=NAVY)
        txt(s, x + 0.22, y + 0.98, 5.7, 0.85, body, size=13, color=INK)

    round_rect(s, 0.45, 5.9, 12.4, 1.15, NAVY, adj=0.08)
    txt(s, 0.7, 6.05, 2.2, 0.28, "CORE INSIGHT", size=11, bold=True, color="9ED4E2")
    txt(s, 0.7, 6.35, 11.9, 0.5,
        "The gap is an AI-first daily OS: gate + bills + home services, in the society's "
        "language — and a preschool mode that never says 'flat' or 'maintenance'.",
        size=15, color=WHITE)
    footer(s, 3)
    return s


def slide_solution(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "What we are building")
    heading(s, 0.5, 0.52, 12, "GATEZO — one operating layer")
    txt(s, 0.5, 1.05, 12.3, 0.45,
        "Mobile + web, one product. Four roles. Two organisation types. "
        "Designed to be used every day, not installed and forgotten.",
        size=14, color=MUTED)

    pillars = [
        (TEAL, "Ask GATEZO",
         "Type or speak. Answers come from live bills, visitors and bookings. Deep-links into the right screen."),
        (NAVY, "Gate & people",
         "Visitor log, gate pass, voice-fill for guards, vehicle QR (Platinum), SOS, directory."),
        (GREEN, "Money",
         "Maintenance / preschool fees, late fees, Razorpay UPI (GPay, PhonePe, Paytm), transparency ledger."),
        (AMBER, "Home & school",
         "On-demand services catalogue, amenities / hall, helpdesk, marketplace, AGM — hidden where a preschool should not see them."),
    ]
    for i, (col, title, body) in enumerate(pillars):
        x = 0.45 + i * 3.2
        round_rect(s, x, 1.6, 3.05, 3.15, WHITE, adj=0.08, line=LINE)
        rect(s, x, 1.6, 3.05, 0.1, col)
        txt(s, x + 0.16, 1.85, 2.75, 0.7, title, size=16, bold=True, color=NAVY)
        txt(s, x + 0.16, 2.55, 2.75, 1.9, body, size=13, color=INK)

    round_rect(s, 0.45, 4.95, 12.4, 2.05, WHITE, adj=0.08, line=LINE)
    txt(s, 0.7, 5.1, 12, 0.3, "What is true today  —  not a mock", size=13, bold=True, color=TEAL)
    bullets = [
        "Working demo for resident, guard, admin and superadmin — live local product and a deployable web build.",
        "Home dashboard: 30-day visitor donut + bill/fee charts; tap through to those tabs.",
        "Dual vocabulary engine: society (flat, maintenance, clubhouse) vs preschool (student, fees, hall, parent).",
        "Commercial packaging already in the product: Base / Prime / Platinum (AI and smart-gate sit on higher tiers).",
    ]
    for i, b in enumerate(bullets):
        txt(s, 0.7, 5.45 + i * 0.35, 11.9, 0.35, "▸  " + b, size=13, color=INK)
    footer(s, 4)
    notes(s, "Emphasise working product, not revenue. Platinum vehicle-gate is in the product as a tier, not a hardware fleet we operate.")
    return s


def slide_tabs(prs, media):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.4, 0.22, "Product  ·  the app, tab by tab")
    heading(s, 0.4, 0.42, 8.2, "What a resident actually opens")
    txt(s, 0.4, 0.92, 8.2, 0.38,
        "Four tabs. Guard and admin use the same product, with the screens that match their job.",
        size=13, color=MUTED)

    phone = media / "phone-home.png"
    if phone.exists():
        pic(s, phone, 0.28, 1.28, 4.05, 5.72)
    else:
        round_rect(s, 0.35, 1.35, 3.9, 5.5, WHITE, adj=0.08, line=LINE)
        txt(s, 0.5, 3.6, 3.6, 0.5, "Home screenshot", size=13, color=MUTED, align="center")

    tabs = [
        ("HOME", TEAL,
         "Ask GATEZO, visitor & bill charts, shortcuts, optional neighbourhood offers."),
        ("MAINTENANCE / FEES", NAVY,
         "Pay by UPI (GPay, PhonePe, Paytm). Receipts, late fees, Transparency score — members see where money went."),
        ("VISITORS", GREEN,
         "Approve guests, gate pass, vehicle QR on Platinum. Guard can speak a visitor in."),
        ("COMMUNITY", AMBER,
         "Helpdesk & SOS, hall / clubhouse, Buy & Sell, photographed home-services catalogue."),
    ]
    for i, (title, col, body) in enumerate(tabs):
        y = 1.28 + i * 1.18
        round_rect(s, 4.5, y, 8.4, 1.08, WHITE, adj=0.1, line=LINE)
        rect(s, 4.5, y, 0.1, 1.08, col)
        txt(s, 4.78, y + 0.1, 7.9, 0.32, title, size=13, bold=True, color=NAVY)
        txt(s, 4.78, y + 0.44, 7.9, 0.52, body, size=13, color=INK)

    round_rect(s, 4.5, 6.05, 8.4, 0.95, NAVY, adj=0.1)
    txt(s, 4.7, 6.18, 8.05, 0.7,
        "Admin / owner: collections, members from a spreadsheet, manager AI, tenants & tiers. "
        "Preschool uses the same tabs with school words (Fees, hall, parent).",
        size=12, color=WHITE)
    footer(s, 6)
    notes(s, "Walk the four tabs in the live app after this slide. Home photo is a real screenshot; Maintenance and Visitors still need live walkthrough until we capture those tabs.")
    return s


def slide_money_tab(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "Product  ·  Maintenance / Fees tab")
    heading(s, 0.5, 0.5, 12, "Bills and school fees, paid in-app")
    txt(s, 0.5, 1.05, 12.2, 0.4,
        "The same payment stack serves a society treasurer and a preschool accounts desk.",
        size=14, color=MUTED)

    items = [
        ("Society", "Maintenance bills, late-fee policy, maintenance heads (sinking fund, water, common). UPI to the society."),
        ("Preschool", "Fees — never labelled “maintenance”. Remind pending fees. Student, not flat."),
        ("How they pay", "Razorpay opens on UPI first: Google Pay, PhonePe, Paytm or a UPI ID. Card and net banking stay available."),
        ("Transparency", "A 0–100 score from the real books: payments in, labelled expenses out, sealed on a hash-chained ledger. Edit history and the chain flags it. Members open this from Bills — not a year-end PDF."),
        ("Admin", "Generate bills, collections dashboard, unpaid list, manager AI for a dues reminder. Seal / re-verify the ledger."),
        ("Why it is cheaper to run", "One integration, not a gate app + Excel + a treasurer’s personal UPI. Base tier includes billing — sites do not buy Platinum to collect fees."),
    ]
    for i, (t, b) in enumerate(items):
        x = 0.45 + (i % 3) * 4.2
        y = 1.6 + (i // 3) * 2.5
        round_rect(s, x, y, 4.0, 2.3, WHITE, adj=0.08, line=LINE)
        txt(s, x + 0.2, y + 0.18, 3.6, 0.4, t, size=15, bold=True, color=TEAL)
        txt(s, x + 0.2, y + 0.65, 3.6, 1.45, b, size=13, color=INK)
    footer(s, 7)
    notes(s, "Demo Transparency after a payment. Do not open Sustainability unless asked — it needs office water-meter readings.")
    return s


def slide_ease_and_features(prs):
    """Catalogue of why it is easy, plus what is actually in the product today."""
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.4, 0.22, "Product  ·  selling points")
    heading(s, 0.4, 0.42, 12.4, "Easy to live in  ·  complete without extra apps")
    txt(s, 0.4, 0.92, 12.4, 0.32,
        "One login. The next action is a sentence, a tap, or a spoken name — not a hunt through five products.",
        size=14, color=MUTED)

    round_rect(s, 0.4, 1.32, 6.25, 5.5, WHITE, adj=0.08, line=LINE)
    txt(s, 0.58, 1.46, 5.95, 0.32, "WHY IT IS EASY", size=13, bold=True, color=TEAL)
    ease = [
        ("Ask, don’t hunt", "Full questions on Home — “How much do I owe?” — answered from live data, with a button into the right screen."),
        ("Charts that open the tab", "Visitor pie and bill chart are not decoration. Tap them and Visitors or Maintenance opens."),
        ("Voice at the gate", "Guard speaks the visitor’s name; the form fills. Residents can talk to Ask GATEZO too."),
        ("Go live from a spreadsheet", "Admin pastes a member / flat list. No billed implementation team to type every household."),
        ("Same UPI everywhere", "One Razorpay checkout for maintenance, school fees and AC service — GPay, PhonePe, Paytm or a UPI ID."),
        ("Help is in the app", "How-to guide rewrites itself for preschool. Parents never see “flat”, “AGM” or “maintenance”."),
        ("The right screen for the job", "Resident, guard, admin, owner. Finance stays off the guard’s desk. Preschool hides society-only modules."),
    ]
    for i, (t, b) in enumerate(ease):
        y = 1.82 + i * 0.68
        txt(s, 0.58, y, 5.95, 0.22, t, size=13, bold=True, color=NAVY)
        txt(s, 0.58, y + 0.22, 5.95, 0.44, b, size=12, color=INK)

    round_rect(s, 6.85, 1.32, 6.05, 5.5, WHITE, adj=0.08, line=LINE)
    txt(s, 7.02, 1.46, 5.7, 0.32, "IN THE PRODUCT TODAY", size=13, bold=True, color=TEAL)
    groups = [
        ("Gate", "Visitor log · approve / deny / leave at gate · gate pass · guard voice log · preschool pickup · vehicle QR (Platinum)."),
        ("Money", "Bills / fees · late fees · maintenance heads · receipts · UPI · collections · Transparency score (0–100, tamper-evident ledger)."),
        ("Community", "Helpdesk · SOS · directory · Buy & Sell · home services catalogue · hall / clubhouse · AGM & voting · trusted helpers · announcements."),
        ("Run the site", "Generate bills · unpaid list · manager AI notices · members from a spreadsheet · Base / Prime / Platinum · superadmin tenants."),
    ]
    for i, (t, b) in enumerate(groups):
        y = 1.86 + i * 1.12
        txt(s, 7.02, y, 5.7, 0.26, t, size=14, bold=True, color=NAVY)
        txt(s, 7.02, y + 0.28, 5.7, 0.78, b, size=12, color=INK)

    footer(s, 8)
    notes(
        s,
        "Do not headline Sustainability. Water / green score exists but needs office meter readings (or a CSV of litres) — extra config, easy to look empty in a demo. "
        "Transparency is live from payments + expenses already in the books — demo that.",
    )
    return s


def slide_gate_tab(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "Product  ·  Visitors tab")
    heading(s, 0.5, 0.5, 12, "The gate, without a paper register")

    cards = [
        ("Resident / parent",
         "See who came. Approve or reject from the phone. Pre-approve a guest with a gate pass. Charts on Home tap through to this tab."),
        ("Guard",
         "Gate desk is the job. Speak the visitor’s name and the form fills. Pending list. No finance screens in the way."),
        ("Preschool pickup",
         "The same gate log, in school language. Pickup at the gate — not a housing-society visitor flow with “flat” on the form."),
        ("Platinum add-on",
         "Vehicle QR / smart-gate hooks when a site needs hardware. Base sites still get a proper digital log — they are not forced into a hardware quote."),
    ]
    for i, (t, b) in enumerate(cards):
        x = 0.45 + (i % 2) * 6.4
        y = 1.25 + (i // 2) * 2.55
        round_rect(s, x, y, 6.2, 2.4, WHITE, adj=0.08, line=LINE)
        txt(s, x + 0.25, y + 0.22, 5.7, 0.4, t, size=16, bold=True, color=NAVY)
        txt(s, x + 0.25, y + 0.75, 5.7, 1.4, b, size=14, color=INK)
    footer(s, 9)
    return s


def slide_product_ai(prs, media):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.45, 0.26, "Product  ·  live screenshot")
    heading(s, 0.45, 0.48, 7.4, "Home is an AI dashboard")
    txt(s, 0.45, 1.0, 7.3, 0.7,
        "Residents land on Ask GATEZO, an attention strip, and charts for visitors and bills. "
        "The pie is not decoration — tap Visitors or Maintenance and the tab opens.",
        size=14, color=MUTED)

    points = [
        ("Ask a full question", "Suggestions are complete sentences — “How much maintenance do I owe?” — answered from live data."),
        ("Voice at the gate", "Guards can speak a visitor in; the form fills. Residents can talk to Ask GATEZO too."),
        ("Role-aware", "Admin asks about collections and spreadsheet import. Guard asks to log a visitor. Preschool hides society-only topics."),
        ("Help built in", "How-to guide rewrites itself for preschool so parents never see “flat” or “AGM”."),
    ]
    for i, (t, b) in enumerate(points):
        y = 1.75 + i * 1.15
        round_rect(s, 0.45, y, 7.35, 1.05, WHITE, adj=0.1, line=LINE)
        txt(s, 0.65, y + 0.12, 7.0, 0.32, t, size=15, bold=True, color=NAVY)
        txt(s, 0.65, y + 0.46, 7.0, 0.48, b, size=13, color=INK)

    phone = media / "phone-home.png"
    if phone.exists():
        # Portrait screenshot; keep aspect, pin to the right.
        pic(s, phone, 8.05, 0.35, 4.85, 6.85)
    footer(s, 5)
    notes(s, "Walk the live home: Ask “how much do I owe?”, then tap the Visitors chart.")
    return s


def slide_roles(prs, media):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "Product  ·  who uses it")
    heading(s, 0.5, 0.5, 12, "Four roles  ·  two organisation types")

    roles = [
        (media / "role-resident.png", "Resident / parent",
         "Pay bills or fees, approve guests, book hall/clubhouse, raise tickets, SOS, Buy & Sell, home services."),
        (media / "role-guard.png", "Guard",
         "Gate desk. Voice-fill visitor. Pending approvals. Preschool pickup flow. No finance clutter."),
        (media / "role-admin.png", "Admin",
         "Collections, bills, Transparency score, members from a spreadsheet, manager AI, gate log."),
        (media / "role-preschool.png", "Preschool mode",
         "Same product, school language: student, fees, hall, parent. Pickup at the gate — never “flat” or “maintenance”."),
    ]
    for i, (img, title, body) in enumerate(roles):
        x = 0.4 + i * 3.22
        round_rect(s, x, 1.2, 3.08, 4.35, WHITE, adj=0.08, line=LINE)
        if img.exists():
            pic(s, img, x, 1.2, 3.08, 1.7)
        txt(s, x + 0.14, 3.05, 2.8, 0.55, title, size=15, bold=True, color=NAVY)
        txt(s, x + 0.14, 3.6, 2.8, 1.7, body, size=12, color=INK)

    round_rect(s, 0.4, 5.7, 12.5, 1.35, NAVY, adj=0.08)
    txt(s, 0.65, 5.85, 12, 0.28, "DUAL VOCABULARY  ·  SUPERADMIN RUNS BOTH", size=12, bold=True, color="9ED4E2")
    txt(s, 0.65, 6.2, 12.1, 0.65,
        "Society: Flat · Maintenance · Clubhouse · Resident.   "
        "Preschool: Student · Fees · Hall · Parent.   "
        "Superadmin: multi-tenant + Base / Prime / Platinum. Home services, AGM and vehicle-gate hide for preschool.",
        size=14, color=WHITE)
    footer(s, 11)
    return s


def slide_services(prs, media):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.45, 0.26, "Product  ·  Community tab")
    heading(s, 0.45, 0.48, 8, "Community market, helpdesk, home services")
    txt(s, 0.45, 1.0, 7.5, 0.6,
        "Residents book electrician, cleaning, AC, salon, movers and more from a photographed catalogue. "
        "Checkout is Razorpay with UPI first — Google Pay, PhonePe, Paytm or a UPI ID.",
        size=14, color=MUTED)

    names = [
        ("instant.jpg", "Instant"),
        ("cleaning.jpg", "Cleaning"),
        ("ac.jpg", "AC"),
        ("movers.jpg", "Movers"),
        ("salon.jpg", "Salon"),
        ("pest.jpg", "Pest"),
    ]
    for i, (fn, label) in enumerate(names):
        x = 0.45 + (i % 3) * 2.5
        y = 1.7 + (i // 3) * 2.15
        img = media / f"svc-{Path(fn).stem}.png"
        round_rect(s, x, y, 2.35, 2.0, WHITE, adj=0.1, line=LINE)
        if img.exists():
            pic(s, img, x, y, 2.35, 1.45)
        txt(s, x + 0.08, y + 1.48, 2.2, 0.42, label, size=13, bold=True, color=NAVY, align="center")

    ac = media / "ac-service.png"
    if ac.exists():
        pic(s, ac, 7.9, 1.55, 5.0, 4.35)
    else:
        round_rect(s, 7.9, 1.55, 5.0, 4.35, WHITE, line=LINE)

    round_rect(s, 7.9, 6.05, 5.0, 0.95, SOFT, adj=0.12)
    txt(s, 8.05, 6.15, 4.7, 0.75,
        "Bills use the same UPI stack. This is the path from SaaS seats to take-rate on jobs — without making ads the product.",
        size=12, color=INK)
    footer(s, 10)
    notes(s, "Show AC Service packages if demoing. Do not invent GMV or take-rate percentages.")
    return s


def slide_how(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "Product / service  ·  how it works")
    heading(s, 0.5, 0.5, 12, "A day on GATEZO")

    steps = [
        ("1", "Arrive",
         "Resident, parent, guard or admin signs in. Org type and tier (Base / Prime / Platinum) shape what they see."),
        ("2", "Ask or tap",
         "Ask GATEZO, or use Home charts. Guards can speak a visitor. Admins open finance or upload a member list."),
        ("3", "Act in one workflow",
         "Approve a guest, pay UPI, book a hall, raise a ticket, or schedule AC service — same session."),
        ("4", "Leave an audit trail",
         "Gate log, receipts, Transparency score, SOS. Superadmin sees every tenant."),
    ]
    for i, (n, title, body) in enumerate(steps):
        x = 0.45 + i * 3.2
        round_rect(s, x, 1.25, 3.05, 3.35, WHITE, adj=0.08, line=LINE)
        round_rect(s, x + 0.18, 1.45, 0.5, 0.5, TEAL, adj=0.3)
        txt(s, x + 0.18, 1.48, 0.5, 0.45, n, size=16, bold=True, color=WHITE, align="center")
        txt(s, x + 0.18, 2.15, 2.7, 0.4, title, size=18, bold=True, color=NAVY)
        txt(s, x + 0.18, 2.65, 2.7, 1.7, body, size=13, color=INK)

    round_rect(s, 0.45, 4.8, 12.4, 2.2, WHITE, adj=0.08, line=LINE)
    txt(s, 0.7, 4.95, 12, 0.3, "PACKAGING ALREADY IN THE PRODUCT", size=12, bold=True, color=TEAL)
    packs = [
        ("Base", "Gate log, bills/fees + UPI, transparency score, marketplace, helpdesk, directory."),
        ("Prime", "AI assistant, gate pass, amenities/hall, reminders, rent, exports."),
        ("Platinum", "Vehicle QR / smart-gate hooks, branding, priority analytics."),
    ]
    for i, (t, b) in enumerate(packs):
        x = 0.7 + i * 4.05
        txt(s, x, 5.35, 3.85, 0.32, t, size=16, bold=True, color=NAVY)
        txt(s, x, 5.7, 3.85, 1.05, b, size=13, color=INK)
    footer(s, 12)
    return s


def slide_why(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.45, 0.26, "Why we can price below the noisy category")
    heading(s, 0.45, 0.48, 12.3, "Cheaper to buy  ·  cheaper to build  ·  better fit")
    txt(s, 0.45, 1.05, 12.3, 0.38,
        "We do not invent a “X% cheaper than MyGate” number. The cost structure is different — and that is the point.",
        size=14, color=MUTED)

    round_rect(s, 0.4, 1.5, 6.3, 5.5, WHITE, adj=0.08, line=LINE)
    txt(s, 0.6, 1.65, 5.95, 0.35, "WHY A SOCIETY / SCHOOL PAYS LESS", size=13, bold=True, color=TEAL)
    cheap = [
        "Base / Prime / Platinum — a site pays for the gate log without buying smart-gate hardware.",
        "One subscription covers gate + bills/fees + preschool language. Incumbents mean a gate app + Excel + Urban Company.",
        "We get paid by subscription, not by packing Home with ads. Neighbourhood offers can sit at the bottom — they are optional, not the P&L.",
        "Admin uploads a member spreadsheet. No billed three-week implementation team on the invoice.",
        "Preschool is the same app, not a second licence and a second vendor.",
    ]
    for i, line in enumerate(cheap):
        txt(s, 0.65, 2.15 + i * 0.9, 5.85, 0.85, "▸  " + line, size=13, color=INK)

    round_rect(s, 6.9, 1.5, 5.95, 5.5, NAVY, adj=0.08)
    txt(s, 7.1, 1.65, 5.55, 0.35, "WHY THE TECH IS LEANER", size=13, bold=True, color="9ED4E2")
    tech = [
        "One React Native / Expo codebase → iOS, Android and web. We do not staff three native teams and pass that into the price.",
        "Dual-vocabulary engine (org type), not a forked “school product”. One codebase, two businesses.",
        "Ask GATEZO reads live tenant data — not a generic chatbot bolted onto a visitor log.",
        "One Razorpay UPI stack for maintenance, fees and home-service checkout.",
        "Role + preschool filters in software. Superadmin multi-tenant and tiers live in the same product we demo.",
    ]
    for i, line in enumerate(tech):
        txt(s, 7.1, 2.15 + i * 0.9, 5.55, 0.85, "▸  " + line, size=13, color=WHITE)
    footer(s, 13)
    notes(s, "Do not quote a rupee discount vs MyGate. Argue structure: tiers, one Expo app, subscription not ads, preschool included. Do not pitch AWS — hosting can move later without rewriting the app.")
    return s


def slide_market(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "Where we play")
    heading(s, 0.5, 0.5, 12, "Market opportunity")
    txt(s, 0.5, 1.05, 12, 0.4,
        "Sizing stays directional until pilot pricing is locked. We will not put a vanity TAM on this slide.",
        size=14, color=MUTED)

    boxes = [
        ("TAM", "India-wide",
         "Gated residential communities + managed workplaces + a second wedge in urban preschools that still run the gate on paper and fees on Excel."),
        ("SAM", "Urban / Pune-led",
         "Digitally adoptable societies and schools where visitors, bills and parent pickup are daily, visible operations."),
        ("SOM", "First 3–5 years",
         "City-by-city density. Win a cluster of sites, then expand. Same playbook for a preschool group as for a society federation."),
    ]
    for i, (k, sub, body) in enumerate(boxes):
        x = 0.45 + i * 4.2
        round_rect(s, x, 1.6, 4.0, 3.15, WHITE, adj=0.08, line=LINE)
        txt(s, x + 0.22, 1.8, 3.55, 0.35, k, size=14, bold=True, color=TEAL)
        txt(s, x + 0.22, 2.2, 3.55, 0.4, sub, size=18, bold=True, color=NAVY)
        txt(s, x + 0.22, 2.75, 3.55, 1.7, body, size=13, color=INK)

    round_rect(s, 0.45, 4.95, 6.15, 2.05, WHITE, adj=0.08, line=LINE)
    txt(s, 0.65, 5.1, 5.8, 0.3, "Why now", size=14, bold=True, color=NAVY)
    txt(s, 0.65, 5.5, 5.8, 1.3,
        "UPI is default. Residents already talk to ChatGPT. Guards have phones. "
        "Committees are tired of paying for ads inside the app they bought for security.",
        size=13, color=INK)

    round_rect(s, 6.8, 4.95, 6.05, 2.05, NAVY, adj=0.08)
    txt(s, 7.0, 5.1, 5.7, 0.3, "Early adopters", size=14, bold=True, color="9ED4E2")
    txt(s, 7.0, 5.5, 5.7, 1.3,
        "Mid-to-large Pune societies with active gates, plus preschools that need pickup control and fee collection without housing-society language.",
        size=13, color=WHITE)
    footer(s, 14)
    return s


def slide_model(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "How we make money")
    heading(s, 0.5, 0.5, 12, "Business model")

    left = [
        ("Who pays",
         "Society / preschool management (site SaaS). Residents pay bills and optional home-service jobs through the app."),
        ("How they pay",
         "Recurring site subscription by tier (Base / Prime / Platinum). Optional modules. Take-rate on fulfilled home-service bookings (rate TBD in pilots)."),
        ("ARPA",
         "Positioned for mass adoption vs. category leaders. Final ARPA after first paid contracts — placeholders on the Ask slide stay blank on purpose."),
    ]
    for i, (t, b) in enumerate(left):
        y = 1.2 + i * 1.5
        round_rect(s, 0.45, y, 6.35, 1.38, WHITE, adj=0.1, line=LINE)
        txt(s, 0.65, y + 0.12, 6.0, 0.32, t, size=15, bold=True, color=TEAL)
        txt(s, 0.65, y + 0.48, 6.0, 0.75, b, size=13, color=INK)

    round_rect(s, 7.0, 1.2, 5.85, 4.5, WHITE, adj=0.08, line=LINE)
    txt(s, 7.2, 1.4, 5.5, 0.35, "Revenue streams", size=16, bold=True, color=NAVY)
    streams = [
        "1.  Site subscription (society or preschool)",
        "2.  Tier upgrades — Prime (AI) / Platinum (smart gate)",
        "3.  Home-services take-rate on booked jobs",
        "4.  Enterprise / multi-site groups",
        "5.  Later: integrations and paid workflows",
    ]
    for i, line in enumerate(streams):
        txt(s, 7.2, 1.95 + i * 0.55, 5.45, 0.5, line, size=14, color=INK)
    txt(s, 7.2, 4.75, 5.45, 0.7,
        "Ads are not how we get paid. Home can show optional neighbourhood offers — they are not the revenue model.",
        size=12, italic=True, color=MUTED)

    round_rect(s, 0.45, 5.75, 12.4, 1.25, SOFT, adj=0.1)
    txt(s, 0.65, 5.9, 12, 0.9,
        "Unit economics: low implementation friction + daily usage (gate, bills, ask) + a second line in home services. "
        "CAC, LTV and gross margin will be measured on the first paid pilots — not projected from a spreadsheet today.",
        size=14, color=INK)
    footer(s, 15)
    return s


def slide_competition(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.4, 0.22, "Why we can win a slice")
    heading(s, 0.4, 0.42, 12, "Competition — named, not 'legacy tools'")

    headers = ["Capability", "GATEZO", "MyGate", "NoBrokerHood", "ADDA", "Urban Co."]
    rows = [
        ["AI on live site data", "Yes", "No", "No", "No", "No"],
        ["Mobile + web, one codebase", "Yes", "Native apps", "Native apps", "Native apps", "Apps"],
        ["Pay by tier (Base / Prime / Platinum)", "Yes", "Bundle / quote", "Bundle", "Bundle", "Per job"],
        ["Ads fund the product", "No", "Heavy", "Ads", "Lighter", "Marketplace"],
        ["In-app home services", "Yes", "Ads / partners", "Some", "No", "Core"],
        ["Society + preschool vocabulary", "Yes", "No", "No", "No", "No"],
        ["See where money went", "Yes", "Limited", "Limited", "Some", "No"],
        ["Go-live setup", "Upload a member list", "On-site / sales", "Sales-led", "Sales-led", "Consumer app"],
    ]
    col_w = [2.45, 2.1, 1.65, 1.8, 1.45, 1.7]
    x0, y0, row_h = 0.4, 1.0, 0.42
    # header
    x = x0
    for i, h in enumerate(headers):
        bgc = TEAL if i == 1 else NAVY
        rect(s, x, y0, col_w[i], row_h, bgc)
        txt(s, x, y0, col_w[i], row_h, h, size=11, bold=True, color=WHITE, align="center", anchor="middle")
        x += col_w[i]
    for r, row in enumerate(rows):
        x = x0
        y = y0 + (r + 1) * row_h
        for c, val in enumerate(row):
            shade = WHITE if r % 2 == 0 else "EEF3F5"
            if c == 1:
                shade = SOFT_G if val in ("Yes", "No") else SOFT
            rect(s, x, y, col_w[c], row_h, shade, line=LINE)
            bold = c == 0 or (c == 1 and val in ("Yes", "No"))
            color = GREEN if (c == 1 and val == "Yes") else INK
            txt(s, x, y, col_w[c], row_h, val, size=11, bold=bold, color=color,
                align="center" if c else "left", anchor="middle")
            x += col_w[c]

    round_rect(s, 0.4, 4.82, 12.5, 2.18, NAVY, adj=0.08)
    txt(s, 0.65, 4.98, 12.1, 0.28, "OUR ADVANTAGE", size=12, bold=True, color="9ED4E2")
    txt(s, 0.65, 5.32, 12.1, 1.45,
        "We do not claim to out-scale MyGate. We win on cost structure: one codebase (mobile + web), "
        "pay-by-tier instead of a bundled quote, subscription instead of an ads P&L, and preschool as the same product — not a second vendor.\n\n"
        "Home can show optional neighbourhood offers. That is not how we intend to get paid.",
        size=14, color=WHITE)
    footer(s, 16)
    notes(s, "If challenged on MyGate: agree they own the gate. Our wedge is ask + bills + services + preschool at a lower cost structure.")
    return s


def slide_roadmap(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "Where this goes")
    heading(s, 0.5, 0.5, 12, "Roadmap & milestones")

    phases = [
        ("TODAY", "Working product",
         "Multi-role demo. Ask GATEZO. Charts. UPI. Home services. Society + preschool labels. Superadmin tiers.",
         TEAL),
        ("+12 MONTHS", "Pilot → repeatable GTM",
         "Pune paying sites. Lock pricing. Measure retention, Ask usage, bill collection, jobs booked. Reference customers.",
         NAVY),
        ("+2 YEARS", "City-level density",
         "Repeat the playbook in selected cities. Preschool groups as a second motion. Keep AI the default surface.",
         NAVY),
        ("+5 YEARS", "Category platform",
         "From the front door to the operating layer for communities and early-years campuses — still simple enough to adopt.",
         NAVY),
    ]
    for i, (when, title, body, col) in enumerate(phases):
        x = 0.4 + i * 3.22
        round_rect(s, x, 1.25, 3.08, 4.35, WHITE, adj=0.08, line=LINE)
        rect(s, x, 1.25, 3.08, 0.12, col)
        txt(s, x + 0.16, 1.55, 2.76, 0.32, when, size=12, bold=True, color=TEAL)
        txt(s, x + 0.16, 1.95, 2.76, 0.85, title, size=18, bold=True, color=NAVY)
        txt(s, x + 0.16, 2.9, 2.76, 2.4, body, size=13, color=INK)

    round_rect(s, 0.4, 5.8, 12.5, 1.2, SOFT, adj=0.1)
    txt(s, 0.65, 5.95, 12.1, 0.85,
        "North star: active sites + weekly Ask / gate / pay usage + retained subscriptions + (later) services GMV. "
        "Pilot success is behaviour change — not logos on a slide.",
        size=14, color=INK)
    footer(s, 17)
    return s


def slide_risks(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "What can go wrong")
    heading(s, 0.5, 0.5, 12, "Key risks & mitigation")

    risks = [
        ("Adoption",
         "Guards and secretaries resist leaving WhatsApp / registers.",
         "Design around the current job-to-be-done. Voice log for guards. Prove minutes saved in the first week of a pilot."),
        ("Incumbents",
         "MyGate / NoBrokerHood can copy AI chrome or bundle services.",
         "Ship dual-vertical vocabulary and Ask-on-live-data first. Stay cheaper and less ad-heavy. Win density in Pune before they notice."),
        ("Pricing",
         "Too cheap to build a company; too expensive vs. 'free + ads'.",
         "Three tiers already in the product. Pilot multiple packages. Charge for Prime (AI) once Ask is habit."),
        ("Brand / category",
         "Nearby names exist in stores (e.g. Homefy Gatemate). Confusion at install.",
         "Legal and store clearance before paid acquisition. This deck and the live demo are branded GATEZO."),
        ("Reliability & privacy",
         "Gate + fees + SOS cannot go down or leak.",
         "Access control, audit logs, transparency ledger. Reliability is a product requirement, not a later phase."),
        ("Services supply",
         "Home-services GMV needs trusted local crews.",
         "Launch catalogue with clear partners; do not promise Urban Company density on day one. Take-rate only on completed jobs."),
    ]
    for i, (t, risk, mit) in enumerate(risks):
        x = 0.4 + (i % 3) * 4.25
        y = 1.15 + (i // 3) * 2.85
        round_rect(s, x, y, 4.1, 2.7, WHITE, adj=0.08, line=LINE)
        txt(s, x + 0.18, y + 0.12, 3.75, 0.32, t.upper(), size=12, bold=True, color=TEAL)
        txt(s, x + 0.18, y + 0.48, 3.75, 0.85, risk, size=13, bold=True, color=NAVY)
        txt(s, x + 0.18, y + 1.35, 3.75, 1.15, mit, size=12, color=INK)
    footer(s, 18)
    return s


def slide_ask(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "The raise")
    heading(s, 0.5, 0.5, 12, "Financials & the ask")

    round_rect(s, 0.45, 1.2, 4.1, 3.55, NAVY, adj=0.08)
    txt(s, 0.7, 1.4, 3.7, 0.3, "WE WANT TO RAISE", size=12, bold=True, color="9ED4E2")
    txt(s, 0.7, 1.85, 3.7, 0.7, "Strategic\npre-seed capital", size=24, bold=True, color=WHITE)
    txt(s, 0.7, 3.0, 3.7, 1.4,
        "Amount: to be finalized against the pilot plan.\n\nUse this round to turn a working demo into paying Pune sites.",
        size=14, color="D5E4EA")

    round_rect(s, 4.75, 1.2, 4.1, 3.55, WHITE, adj=0.08, line=LINE)
    txt(s, 4.95, 1.4, 3.7, 0.3, "USE OF FUNDS", size=12, bold=True, color=TEAL)
    txt(s, 4.95, 1.9, 3.7, 2.5,
        "[X]%   →   Product / tech\n"
        "[X]%   →   Marketing / GTM\n"
        "[X]%   →   Team / salaries\n"
        "[X]%   →   Operations",
        size=16, color=INK)

    round_rect(s, 9.05, 1.2, 3.85, 3.55, WHITE, adj=0.08, line=LINE)
    txt(s, 9.25, 1.4, 3.5, 0.3, "3-YEAR SNAPSHOT", size=12, bold=True, color=TEAL)
    for i, yr in enumerate(("Year 1", "Year 2", "Year 3")):
        y = 1.9 + i * 0.85
        txt(s, 9.25, y, 3.5, 0.28, yr, size=13, bold=True, color=MUTED)
        txt(s, 9.25, y + 0.28, 3.5, 0.4, "[ ₹ ]", size=22, bold=True, color=NAVY)

    round_rect(s, 0.45, 4.95, 12.4, 2.05, WHITE, adj=0.08, line=LINE)
    txt(s, 0.7, 5.1, 12, 0.3, "BEYOND CAPITAL", size=12, bold=True, color=TEAL)
    txt(s, 0.7, 5.45, 12, 1.3,
        "Introductions to society federations, preschool groups, facility leaders, proptech mentors and early pilot customers.\n"
        "That support unlocks: first paying sites, pricing truth, and a Pune reference set we can take to the next city.\n\n"
        "Investor model (ARR, burn, runway) will be filled after pilot pricing, conversion, retention and site-level costs are real.",
        size=14, color=INK)
    footer(s, 19)
    notes(s, "Leave [X]% and [ ₹ ] as blanks unless Udit supplies numbers. Do not improvise a raise amount.")
    return s


def slide_annex_gtm(prs):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    kicker(s, 0.5, 0.28, "Annexure 1")
    heading(s, 0.5, 0.5, 12, "Go-to-market  ·  customer value")

    phases = [
        ("Phase 1 — Pune pilots",
         "High-density societies + 1–2 preschools.\nFounder-led sales and implementation.\nCapture usage, objections, retention, pricing."),
        ("Phase 2 — Playbook",
         "Standard onboarding (admin uploads a member spreadsheet).\nReference customers.\nChannel / secretary / CA referrals."),
        ("Phase 3 — Cities",
         "Replicate proven economics.\nPreschool groups as a parallel motion.\nKeep Ask GATEZO the first surface."),
    ]
    for i, (t, b) in enumerate(phases):
        x = 0.45 + i * 4.2
        round_rect(s, x, 1.2, 4.0, 2.55, WHITE, adj=0.08, line=LINE)
        txt(s, x + 0.2, 1.35, 3.6, 0.55, t, size=15, bold=True, color=NAVY)
        txt(s, x + 0.2, 1.95, 3.6, 1.55, b, size=13, color=INK)

    value = [
        ("Residents / parents", "Ask instead of hunting. Pay UPI. See visitor charts. Book a plumber. SOS."),
        ("Guards", "Voice log. Clear pending list. Preschool pickup without a housing-society UI."),
        ("Management", "Collections, members from a spreadsheet, notices via manager AI, transparency ledger."),
        ("The company", "Daily-use SaaS + tier upgrades + services take-rate. Retention from the front door."),
    ]
    for i, (t, b) in enumerate(value):
        x = 0.45 + (i % 2) * 6.4
        y = 3.95 + (i // 2) * 1.4
        round_rect(s, x, y, 6.2, 1.28, WHITE, adj=0.1, line=LINE)
        txt(s, x + 0.2, y + 0.12, 5.8, 0.32, t, size=14, bold=True, color=TEAL)
        txt(s, x + 0.2, y + 0.48, 5.8, 0.65, b, size=13, color=INK)
    footer(s, 20)
    return s


def slide_close(prs, media):
    s = blank(prs)
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, W, 0.12, TEAL)
    wordmark = ASSETS / "brand" / "gatezo" / "wordmark.png"
    if wordmark.exists():
        pic(s, wordmark, 0.45, 0.22, 3.4, 0.55)
        kicker(s, 4.05, 0.32, "Annexure 2  ·  KPIs  ·  thesis")
        heading(s, 0.5, 0.82, 12, "What we will measure  ·  why it matters")
    else:
        kicker(s, 0.5, 0.26, "Annexure 2  ·  KPIs  ·  thesis")
        heading(s, 0.5, 0.48, 12, "What we will measure  ·  why it matters")

    kpis = [
        "Active sites (societies + preschools)",
        "Weekly active users by role",
        "Ask GATEZO questions / user / week",
        "Gate events logged vs. paper/WhatsApp",
        "% bills collected via UPI",
        "Home-service bookings (when live)",
        "Logo retention / logo churn",
        "Revenue per site  ·  CAC payback  ·  gross margin",
    ]
    round_rect(s, 0.45, 1.15, 6.3, 4.35, WHITE, adj=0.08, line=LINE)
    txt(s, 0.65, 1.3, 6.0, 0.35, "North-star metrics", size=16, bold=True, color=NAVY)
    for i, k in enumerate(kpis):
        txt(s, 0.7, 1.75 + i * 0.42, 5.85, 0.4, "▸  " + k, size=13, color=INK)

    round_rect(s, 6.95, 1.15, 5.9, 4.35, NAVY, adj=0.08)
    txt(s, 7.15, 1.35, 5.5, 0.35, "INVESTOR THESIS", size=13, bold=True, color="9ED4E2")
    txt(s, 7.15, 1.85, 5.5, 3.35,
        "The gate is still the daily touchpoint.\n\n"
        "The wedge is no longer 'make entry easier' alone — it is Ask GATEZO on live community data, then bills and home services in the same session.\n\n"
        "The expansion is a trusted OS for societies and preschools, site by site, city by city.\n\n"
        "We are pre-revenue. The product is demo-ready.",
        size=14, color=WHITE)

    round_rect(s, 0.45, 5.65, 12.4, 1.35, WHITE, adj=0.1, line=LINE)
    txt(s, 0.7, 5.8, 12, 0.28, "LIVE DEMO WALK (10 MINUTES)", size=12, bold=True, color=TEAL)
    txt(s, 0.7, 6.15, 12, 0.65,
        "1) Resident home — Ask “what do I owe?”  2) Tap the Visitors chart  3) AC Service packages  "
        "4) Guard voice log  5) Admin collections / member list  6) Superadmin tenants & tiers.",
        size=14, color=INK)
    footer(s, 21)
    notes(s, "Close on thesis + offer the live demo. Credentials stay off the slide.")
    return s


def write_outputs(prs):
    OUT.parent.mkdir(parents=True, exist_ok=True)
    dest = OUT
    try:
        prs.save(dest)
    except PermissionError:
        dest = OUT_ALT
        prs.save(dest)
        print(f"Latest file is open — wrote {dest} instead")
    demo = DECK / "GATEZO_Demo_Deck.pptx"
    try:
        prs.save(demo)
        print(f"Wrote {demo}  ({demo.stat().st_size // 1024} KB)")
    except PermissionError:
        print(f"Could not write {demo} (file open)")
    print(f"Wrote {dest}  ({dest.stat().st_size // 1024} KB)")


def main():
    media = prepare_media()
    prs = new_prs()
    slide_title(prs, media)
    slide_team(prs)
    slide_problem(prs)
    slide_solution(prs)
    slide_product_ai(prs, media)
    slide_tabs(prs, media)
    slide_money_tab(prs)
    slide_ease_and_features(prs)
    slide_gate_tab(prs)
    slide_services(prs, media)
    slide_roles(prs, media)
    slide_how(prs)
    slide_why(prs)
    slide_market(prs)
    slide_model(prs)
    slide_competition(prs)
    slide_roadmap(prs)
    slide_risks(prs)
    slide_ask(prs)
    slide_annex_gtm(prs)
    slide_close(prs, media)
    write_outputs(prs)


if __name__ == "__main__":
    main()
